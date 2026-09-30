from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\output\documentos\Tesis_JosephKleynMamaniPerez_capitulo4_actualizado.docx"
)


PAGES = [
    [
        ("CAPÍTULO I", "8", 0),
        ("1.1 Razón social", "8", 1),
        ("1.2 Misión, visión, objetivos y valores de la empresa", "9", 1),
        ("1.3 Productos, mercado y clientes", "10", 1),
        ("1.4 Estructura de la organización", "11", 1),
        ("1.5 Otra información relevante de la empresa donde se desarrolla el proyecto", "12", 1),
        ("CAPÍTULO II", "13", 0),
        ("2.1 Identificación del problema técnico en la empresa", "13", 1),
        ("2.2 Objetivos del proyecto de mejora", "16", 1),
        ("2.3 Antecedentes del proyecto de mejora", "17", 1),
        ("Justificación del proyecto de mejora", "20", 1),
        ("2.4 Marco teórico y conceptual", "22", 1),
        ("2.4.1 Fundamento teórico del proyecto de mejora", "22", 2),
        ("2.4.2 Conceptos y términos utilizados", "31", 2),
    ],
    [
        ("CAPÍTULO III", "39", 0),
        ("3.1 Diagrama del proceso, mapa del flujo de valor y/o diagrama de operación actual", "39", 1),
        ("3.2 Efectos del problema en el área de trabajo o en los resultados de la empresa", "44", 1),
        ("3.3 Análisis de las causas raíz que generan el problema", "46", 1),
        ("3.4 Priorización de causas raíz", "47", 1),
        ("CAPÍTULO IV", "50", 0),
        ("4.1 Plan de acción de la mejora propuesta", "50", 1),
        ("4.2 Consideraciones técnicas, operativas y ambientales para la implementación de la mejora", "51", 1),
        ("4.3 Recursos técnicos para implementar la mejora propuesta", "53", 1),
        ("4.4 Diagrama del proceso, mapa del flujo de valor y diagrama de operación de la situación mejorada", "55", 1),
        ("4.5 Cronograma de ejecución de la mejora", "58", 1),
        ("4.6 Aspectos limitantes para la implementación de la mejora", "59", 1),
    ],
    [
        ("CAPÍTULO V", "61", 0),
        ("5.1 Costo de materiales", "61", 1),
        ("5.2 Costo de mano de obra", "61", 1),
        ("5.3 Costo de máquinas, herramientas y equipos", "61", 1),
        ("5.4 Otros costos de implementación", "61", 1),
        ("5.5 Costo total de la implementación de la mejora", "61", 1),
        ("CAPÍTULO VI", "62", 0),
        ("6.1 Beneficio técnico y/o económico esperado", "62", 1),
        ("6.2 Relación beneficio/costo", "62", 1),
        ("CAPÍTULO VII", "63", 0),
        ("7.1 Conclusiones respecto a los objetivos del proyecto de mejora", "63", 1),
        ("CAPÍTULO VIII", "64", 0),
        ("8.1 Recomendaciones para la empresa respecto del proyecto de mejora", "64", 1),
        ("REFERENCIAS BIBLIOGRÁFICAS", "65", 0),
        ("ANEXOS", "70", 0),
    ],
]


def format_entry(paragraph, title: str, page: str, level: int, page_break_before: bool) -> None:
    paragraph.style = "Normal"
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.18 * level)
    fmt.first_line_indent = Inches(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.0
    fmt.keep_together = True
    fmt.page_break_before = page_break_before
    fmt.tab_stops.add_tab_stop(Inches(6.05 - 0.18 * level), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    run = paragraph.add_run(f"{title}\t{page}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "Times New Roman")


def main() -> None:
    document = Document(DOCX)
    body = document._element.body
    stale_toc = next((element for element in body if element.tag == qn("w:sdt")), None)
    if stale_toc is None:
        raise RuntimeError("No se encontró el bloque de índice obsoleto.")

    insertion_index = body.index(stale_toc)
    body.remove(stale_toc)

    created = []
    for page_index, entries in enumerate(PAGES):
        for entry_index, (title, page, level) in enumerate(entries):
            paragraph = document.add_paragraph()
            starts_new_page = entry_index == 0 and page_index > 0
            format_entry(paragraph, title, page, level, starts_new_page)
            created.append(paragraph._p)

    for offset, element in enumerate(created):
        body.insert(insertion_index + offset, element)

    document.save(DOCX)
    print(f"Índice actualizado: {DOCX}")


if __name__ == "__main__":
    main()
