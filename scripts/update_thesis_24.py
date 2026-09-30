from __future__ import annotations

import os
import shutil
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"C:\Users\Power\Downloads\Tesis_JosephKleynMamaniPerez_capitulo3.docx")
OUTPUT = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\output\documentos"
    r"\Tesis_JosephKleynMamaniPerez_capitulo3_actualizado.docx"
)


THEORY_SECTIONS: list[tuple[str, list[str]]] = [
    (
        "Fundamentos del sistema de gestión para la tienda",
        [
            "El proyecto se sustenta en la digitalización y centralización de los procesos de Minimarket Tienda Perez. El sistema reúne en una sola aplicación la información de productos, existencias, proveedores, compras, percepciones, ventas, caja, cuentas por cobrar, clientes frecuentes y seguridad. Esta integración permite sustituir registros dispersos en cuadernos por datos estructurados, consultables y trazables, conservando el alcance tributario definido para el negocio.",
            "El sistema no tiene como finalidad emitir boletas ni facturas electrónicas. Los comprobantes de compra continuarán siendo emitidos por los proveedores y serán registrados para organizar las adquisiciones, los pagos y las percepciones. En las ventas se generará un ticket interno con el detalle y el importe calculado; este documento servirá como constancia operativa y no reemplazará un comprobante de pago tributario. Además, los límites de ingresos y adquisiciones del Nuevo RUS podrán mostrarse como información de control para apoyar al propietario, de acuerdo con los parámetros publicados por la Superintendencia Nacional de Aduanas y de Administración Tributaria (SUNAT, s. f.-a, s. f.-b).",
        ],
    ),
    (
        "Arquitectura cliente-servidor y organización modular",
        [
            "La aplicación adopta una arquitectura cliente-servidor: el navegador presenta la interfaz y envía solicitudes, mientras que el servidor valida los datos, aplica las reglas del negocio y accede a la base de datos. La comunicación principal utiliza HTTP y una API orientada a recursos; los mensajes se intercambian en formato JSON, definido como un formato textual y liviano para representar datos estructurados (Fielding, 2000; Bray, 2017).",
            "El backend se organiza mediante módulos con responsabilidades delimitadas, por ejemplo: autenticación, usuarios, sucursales, productos, almacenes, inventario, compras, ventas, caja, deudas, clientes y reportes. NestJS utiliza módulos, controladores y proveedores con inyección de dependencias, lo que facilita aislar las reglas de cada dominio, realizar pruebas y ampliar el sistema sin modificar toda la aplicación (NestJS, s. f.-a).",
        ],
    ),
    (
        "Frontend con Vue.js y TypeScript",
        [
            "La interfaz se desarrolla con Vue.js y TypeScript. Vue.js emplea un modelo declarativo y basado en componentes, mientras que TypeScript incorpora comprobación estática de tipos sobre JavaScript. Esta combinación permite construir pantallas reutilizables para ventas, compras, inventario, caja y clientes, y detectar inconsistencias de datos durante el desarrollo (Vue.js, s. f.; Microsoft, s. f.).",
            "Vuetify aporta componentes de interfaz y un diseño adaptable; Vue Router administra la navegación de la aplicación; y Pinia mantiene el estado compartido, como la sesión, los permisos y la sucursal activa. Axios se utiliza para realizar solicitudes HTTP, Zod valida estructuras de datos y Vite proporciona el servidor de desarrollo y la construcción optimizada para producción (Vuetify, s. f.; Vue Router, s. f.; Pinia, s. f.; Axios, s. f.; Zod, s. f.; Vite, s. f.).",
        ],
    ),
    (
        "Backend con Node.js, TypeScript y NestJS",
        [
            "Node.js constituye el entorno de ejecución del servidor. Su modelo de entrada y salida asíncrona resulta adecuado para atender solicitudes de red, consultas a la base de datos y generación de documentos sin bloquear innecesariamente otras operaciones. Sobre esta plataforma, NestJS proporciona una estructura para desarrollar aplicaciones del lado del servidor con TypeScript (Node.js, s. f.; NestJS, s. f.-a).",
            "Los controladores exponen las operaciones de la API y los servicios concentran las reglas del negocio. Los objetos de transferencia de datos y los esquemas de validación delimitan la información aceptada por cada operación. Esta separación es necesaria para comprobar cantidades, precios, existencias, estados de caja, saldos de deuda y relaciones entre compras, ventas y movimientos de inventario antes de guardar los cambios.",
        ],
    ),
    (
        "API REST, JSON y documentación de servicios",
        [
            "La API utiliza recursos y métodos HTTP para consultar, crear, actualizar o anular registros. Las respuestas en JSON permiten que el frontend y otros servicios interpreten la información de manera uniforme. El enfoque REST promueve una interfaz uniforme y una comunicación sin estado entre cliente y servidor, mientras que el estándar OpenAPI permite describir rutas, parámetros y respuestas de una API (Fielding, 2000; Bray, 2017; OpenAPI Initiative, 2021).",
            "La integración de Swagger en NestJS permitirá documentar los endpoints y facilitar las pruebas durante el desarrollo. Esta documentación será especialmente útil para los módulos de compras, inventario, ventas, caja y para la futura comunicación con el servicio de análisis de video.",
        ],
    ),
    (
        "Persistencia con MySQL y Prisma ORM",
        [
            "MySQL almacena la información en tablas relacionadas y permite establecer claves, restricciones e integridad referencial. El modelo de datos representa entidades como producto, categoría, proveedor, cliente, compra, detalle de compra, venta, detalle de venta, almacén, movimiento de inventario, caja, deuda, abono y comprobante recibido (Oracle, 2026).",
            "Prisma ORM conecta el backend con MySQL mediante un esquema declarativo y consultas tipadas. Prisma Client facilita las operaciones de lectura y escritura, mientras que las migraciones permiten mantener la evolución de la estructura de la base de datos. Las transacciones deben agrupar cambios dependientes; por ejemplo, registrar una compra y aumentar el stock, o registrar una venta y disminuirlo, evitando que una operación quede aplicada solo de forma parcial (Prisma, s. f.).",
        ],
    ),
    (
        "Seguridad, autenticación y autorización",
        [
            "La autenticación verifica la identidad del usuario y la autorización determina las acciones que puede realizar. El sistema emplea JWT para representar datos de sesión, cookies con opciones de seguridad, Passport para integrar estrategias de autenticación y bcrypt para proteger las contraseñas mediante hash. Las contraseñas nunca deben almacenarse como texto sin protección (Jones et al., 2015; OWASP Foundation, s. f.-a; OWASP Foundation, s. f.-b).",
            "El control de acceso se basa en roles y permisos. De este modo, un administrador, vendedor u otro responsable solo accede a los módulos y operaciones que le corresponden. CASL permite expresar capacidades y restricciones en la aplicación, mientras que Helmet, el control de CORS y la limitación de solicitudes complementan la protección del servidor. El modelo RBAC simplifica la administración al asociar permisos con roles y luego asignar dichos roles a los usuarios (NIST, s. f.; CASL, s. f.).",
        ],
    ),
    (
        "Comunicación en tiempo real mediante WebSocket y Socket.IO",
        [
            "WebSocket establece un canal bidireccional persistente entre el navegador y el servidor, evitando crear una nueva conexión HTTP para cada actualización. Socket.IO incorpora comunicación basada en eventos, reconexión automática y mecanismos alternativos de transporte cuando WebSocket no está disponible (Fette & Melnikov, 2011; Socket.IO, s. f.).",
            "NestJS integra Socket.IO mediante gateways. En el proyecto, los eventos en tiempo real se utilizarán para notificar cambios de stock, nuevas ventas, movimientos de caja, actualización de deudas y alertas del módulo de videovigilancia. El frontend deberá autenticarse también al conectarse al socket y recibir únicamente los eventos autorizados para su usuario o sucursal (NestJS, s. f.-b).",
        ],
    ),
    (
        "Gestión de productos, inventario y kardex",
        [
            "El inventario comprende los bienes destinados a la venta. Su control requiere conocer las entradas, salidas, ajustes y existencias disponibles de cada producto. La NIC 2 establece principios para el tratamiento de los inventarios, mientras que el sistema aplicará un kardex digital que conserve la secuencia de movimientos y el saldo resultante por producto y almacén (IFRS Foundation, 2021).",
            "Las compras incrementan las existencias y las ventas las disminuyen; los ajustes corrigen diferencias justificadas. Cada movimiento guardará fecha, cantidad, tipo, documento relacionado y usuario responsable. Con esta trazabilidad se podrán identificar productos con bajo stock, revisar reposiciones y detectar inconsistencias entre el registro y la mercadería física.",
        ],
    ),
    (
        "Compras, proveedores, gastos y percepciones",
        [
            "El módulo de compras registrará al proveedor, los productos recibidos, cantidades, costos, forma de pago y datos del comprobante emitido por el proveedor. También permitirá adjuntar o consignar la información necesaria para localizar el documento físico o digital. Los pagos parciales y pendientes se relacionarán con la compra para conocer el estado de la obligación.",
            "Cuando corresponda, la percepción se registrará como un importe adicional asociado a la operación y al comprobante de percepción recibido. La SUNAT señala que el agente de percepción cobra ese importe al cliente y emite el comprobante correspondiente. El sistema organizará dicha información para consulta y reporte, pero no calculará obligaciones tributarias ni sustituirá la validación del propietario o de un profesional contable (SUNAT, s. f.-b).",
        ],
    ),
    (
        "Ventas, tickets internos y control de caja",
        [
            "El módulo de ventas calculará automáticamente subtotales, descuentos y total, validará las cantidades disponibles y registrará el medio de pago. Al confirmar la operación, el sistema actualizará el inventario y la caja dentro de una transacción. Esta automatización reduce los errores que actualmente pueden producirse al sumar importes y elaborar tickets con papel y lapicero.",
            "El ticket generado será un documento interno con fecha, correlativo, productos, cantidades, precios y monto total. No se presentará como boleta ni factura y no reemplazará los comprobantes exigidos por la normativa aplicable. El cierre de caja consolidará ingresos, egresos, pagos y diferencias para facilitar la revisión diaria.",
        ],
    ),
    (
        "Ventas al crédito, deudas y abonos",
        [
            "Las ventas al crédito o fiado originan una cuenta por cobrar al cliente. El sistema registrará el importe inicial, los abonos, fechas, observaciones y saldo pendiente, manteniendo el historial de cada movimiento. Una regla impedirá que el total abonado exceda la deuda y permitirá identificar obligaciones vencidas o pendientes.",
            "La digitalización de este proceso reemplaza el cuaderno de fiados y mejora la trazabilidad. Sin embargo, la autorización para conceder crédito seguirá siendo una decisión del propietario; el sistema servirá para registrar y consultar la información, no para aprobar automáticamente a un cliente.",
        ],
    ),
    (
        "Clientes frecuentes, promociones y recompensas",
        [
            "El registro voluntario de clientes permitirá asociar sus compras y obtener indicadores de frecuencia, cantidad o monto acumulado. Con reglas configurables podrán establecerse promociones, puntos o recompensas. Las plataformas comerciales de referencia muestran que los programas de fidelización pueden asignar beneficios según las compras y mantener un historial para personalizar la atención (Loyverse, s. f.; Odoo, s. f.).",
            "El tratamiento de los datos personales deberá limitarse a la información necesaria para el programa y contar con la aceptación del cliente. Los reportes deberán mostrar datos pertinentes al usuario autorizado y evitar exponer información personal sin una finalidad operativa.",
        ],
    ),
    (
        "Generación de PDF, Excel y códigos de barras",
        [
            "pdfmake permitirá construir tickets y reportes en formato PDF mediante definiciones de contenido, tablas y estilos. ExcelJS permitirá generar archivos XLSX con hojas, filas, fórmulas y formatos para exportar reportes de compras, ventas, inventario, caja, deudas y clientes frecuentes. Estos documentos serán salidas del sistema y no modificarán la información almacenada (pdfmake, s. f.; ExcelJS, s. f.).",
            "La biblioteca bwip-js podrá generar códigos de barras para etiquetas de productos o para representar identificadores internos. El código de barras agiliza la búsqueda, pero debe vincularse con un producto válido y no sustituye los controles de precio, unidad de medida o existencia (bwip-js, s. f.).",
        ],
    ),
    (
        "Integración de cámaras y transmisión de video",
        [
            "La visualización de cámaras IP se plantea mediante las direcciones y credenciales autorizadas del establecimiento. RTSP es un protocolo de capa de aplicación utilizado para configurar y controlar la entrega de contenido multimedia en tiempo real. OpenCV puede capturar video desde cámaras o flujos y procesar sus fotogramas para el análisis posterior (Schulzrinne et al., 2016; OpenCV, s. f.).",
            "Por seguridad, las credenciales de las cámaras no deben enviarse directamente al navegador ni quedar visibles en el código del frontend. El servidor o un servicio de video intermedio administrará la conexión, controlará el acceso y limitará la conservación de imágenes según la finalidad del proyecto.",
        ],
    ),
    (
        "Servicio de inteligencia artificial para acciones sospechosas",
        [
            "El análisis de video se implementará como un servicio independiente en Python. FastAPI expondrá los endpoints de configuración y resultados; Pydantic validará los datos intercambiados; Uvicorn ejecutará el servicio; OpenCV obtendrá y preparará los fotogramas; y PyTorch proporcionará los componentes para entrenar y ejecutar redes neuronales (Python Software Foundation, s. f.; FastAPI, s. f.; Pydantic, s. f.; Uvicorn, s. f.; OpenCV, s. f.; PyTorch, s. f.).",
            "La separación del servicio de inteligencia artificial evita mezclar el procesamiento intensivo de video con las operaciones transaccionales de la tienda. NestJS recibirá los eventos generados por dicho servicio, guardará la evidencia autorizada y enviará la alerta a los usuarios conectados mediante Socket.IO.",
        ],
    ),
    (
        "Modelo propuesto: YOLOv8, ByteTrack y red 3D CNN",
        [
            "El modelo propuesto se organiza en tres etapas. Primero, YOLOv8 detectará personas y objetos relevantes en cada fotograma. Segundo, ByteTrack asociará las detecciones a lo largo del tiempo para conservar la trayectoria de cada persona. Tercero, una red neuronal convolucional tridimensional analizará secuencias breves de video para clasificar patrones temporales como normales o potencialmente sospechosos (Ultralytics, s. f.; Zhang et al., 2022; Martínez-Mascorro et al., 2020).",
            "La red 3D CNN deberá entrenarse o ajustarse con videos representativos del entorno real de la tienda, considerando diferentes ángulos, iluminación, cantidad de personas y oclusiones. La salida no será una acusación de hurto, sino una puntuación de riesgo y una alerta que deberá ser revisada por una persona. Esta precaución es necesaria porque los falsos positivos pueden afectar injustamente a clientes o trabajadores.",
        ],
    ),
    (
        "Entrenamiento, validación y métricas del modelo",
        [
            "Los videos se dividirán en conjuntos de entrenamiento, validación y prueba sin repetir una misma secuencia entre conjuntos. La evaluación considerará precisión, exhaustividad o recall, F1, matriz de confusión y cantidad de falsas alarmas por periodo. Para la detección también podrá emplearse la intersección sobre unión (IoU), mientras que para el seguimiento se revisará la conservación de identidades (Ultralytics, s. f.; Zhang et al., 2022).",
            "La investigación de Martínez-Mascorro et al. (2020) mostró que una 3D CNN puede identificar segmentos previos a posibles hurtos, pero también evidencia que el resultado depende del conjunto de datos y no elimina la necesidad de supervisión humana. Por ello, el sistema permitirá configurar umbrales, conservar el fragmento relacionado con la alerta y registrar si el operador la confirmó o descartó, datos que servirán para mejorar el modelo.",
        ],
    ),
    (
        "Pruebas, calidad y mantenimiento",
        [
            "Vitest se empleará en las pruebas unitarias y de integración del backend, y Supertest permitirá comprobar los endpoints HTTP. También se realizarán pruebas de los flujos críticos: compra con actualización de stock, venta con caja, abono de deuda, generación de archivos, autorización por roles y entrega de eventos en tiempo real (Vitest, s. f.; SuperAgent Project, s. f.).",
            "Oxlint ayudará a detectar problemas estáticos y Prettier mantendrá un formato uniforme. En el módulo de inteligencia artificial se añadirán pruebas sobre videos separados del entrenamiento y mediciones de rendimiento para comprobar que el procesamiento pueda ejecutarse con el equipo disponible. La validación final comparará los procedimientos digitalizados con los registros manuales actuales.",
        ],
    ),
    (
        "Despliegue y operación del sistema",
        [
            "Durante el desarrollo, Vite servirá la interfaz y NestJS atenderá la API desde procesos separados. Para producción, la interfaz podrá compilarse como archivos estáticos y publicarse junto con el backend o detrás de un servidor web. La base de datos MySQL mantendrá la información operativa, mientras que el servicio de inteligencia artificial se ejecutará de forma independiente para permitir su actualización sin interrumpir las ventas.",
            "La operación deberá incluir copias de seguridad, control de variables de entorno, registro de errores y revisión periódica de permisos. La disponibilidad de cámaras o del modelo de inteligencia artificial no debe bloquear la venta, el inventario ni la caja; si ese servicio falla, el sistema principal continuará funcionando y mostrará el estado de la integración.",
        ],
    ),
]


GLOSSARY: list[tuple[str, str]] = [
    ("API REST", "Interfaz que expone recursos mediante operaciones HTTP y aplica restricciones como una interfaz uniforme y comunicación sin estado. En el proyecto conecta el frontend, el backend y los servicios complementarios (Fielding, 2000)."),
    ("Arquitectura cliente-servidor", "Modelo en el que un cliente solicita recursos o servicios y un servidor procesa la solicitud y devuelve una respuesta. El navegador actúa como cliente y NestJS como servidor (MDN Web Docs, s. f.)."),
    ("Arquitectura modular", "Organización del software en unidades con responsabilidades delimitadas. NestJS emplea módulos para agrupar controladores y proveedores relacionados con un dominio funcional (NestJS, s. f.-a)."),
    ("Autenticación", "Proceso de verificar que una entidad es quien afirma ser mediante uno o más autenticadores. En el sistema antecede a la autorización de cada operación (OWASP Foundation, s. f.-a)."),
    ("Backend", "Parte del sistema que se ejecuta en el servidor, aplica reglas del negocio y coordina el acceso a los datos y servicios. En el proyecto se desarrolla con Node.js, TypeScript y NestJS (Node.js, s. f.; NestJS, s. f.-a)."),
    ("Base de datos relacional", "Colección de datos organizada en tablas vinculadas mediante relaciones y restricciones. MySQL utiliza este modelo para conservar la información estructurada de la tienda (Oracle, 2026)."),
    ("ByteTrack", "Algoritmo de seguimiento multiobjeto que asocia detecciones de alta y baja puntuación para conservar trayectorias e identidades entre fotogramas (Zhang et al., 2022)."),
    ("Cámara IP", "Dispositivo de video conectado a una red que puede proporcionar un flujo digital para su visualización o procesamiento. El acceso a sus transmisiones puede controlarse mediante protocolos como RTSP (Schulzrinne et al., 2016)."),
    ("Cliente frecuente", "Cliente identificado voluntariamente cuyo historial de compras permite medir recurrencia o acumulación y aplicar beneficios definidos por el negocio (Loyverse, s. f.; Odoo, s. f.)."),
    ("Código de barras", "Representación gráfica de datos que puede vincular un identificador con un producto. bwip-js permite generar imágenes de múltiples simbologías de códigos de barras (bwip-js, s. f.)."),
    ("Compra", "Adquisición de productos destinados a la venta. En el sistema origina el registro del proveedor, costo, comprobante recibido, pago y entrada de inventario (IFRS Foundation, 2021)."),
    ("Comprobante de percepción", "Documento emitido por el agente de percepción para acreditar el importe adicional cobrado al cliente dentro del régimen de percepciones del IGV (SUNAT, s. f.-b)."),
    ("Control de acceso basado en roles (RBAC)", "Modelo que asigna permisos a roles y roles a usuarios, simplificando la administración de privilegios según las responsabilidades de la organización (NIST, s. f.)."),
    ("CORS", "Mecanismo basado en cabeceras HTTP mediante el cual un servidor indica desde qué orígenes permite que el navegador acceda a sus recursos (MDN Web Docs, s. f.)."),
    ("Cuenta por cobrar", "Derecho de cobro originado por una venta al crédito. En el sistema se representa mediante la deuda, los abonos y el saldo pendiente del cliente (IFRS Foundation, 2018)."),
    ("Detección de objetos", "Tarea de visión artificial que localiza y clasifica objetos dentro de una imagen o fotograma mediante cajas delimitadoras y puntuaciones de confianza (Ultralytics, s. f.)."),
    ("DTO", "Objeto de transferencia de datos utilizado para definir la forma de la información que entra o sale de una operación. En NestJS puede combinarse con validación para proteger los límites de la API (NestJS, s. f.-a)."),
    ("ExcelJS", "Biblioteca de JavaScript para leer, manipular y generar libros en formatos como XLSX. Se utilizará para exportar reportes tabulares (ExcelJS, s. f.)."),
    ("Falso positivo", "Caso en el que el modelo genera una alerta positiva para una situación que en realidad no pertenece a la clase sospechosa. Debe medirse y revisarse durante la validación (Martínez-Mascorro et al., 2020)."),
    ("FastAPI", "Framework web para construir API con Python a partir de anotaciones de tipo, validación y documentación basada en estándares abiertos (FastAPI, s. f.)."),
    ("Frontend", "Parte de la aplicación que se ejecuta en el navegador y presenta pantallas, formularios y resultados al usuario. En el proyecto se desarrolla con Vue.js y Vuetify (Vue.js, s. f.; Vuetify, s. f.)."),
    ("Gasto", "Disminución de beneficios económicos producida durante un periodo. En el sistema corresponde a una salida registrada y clasificada para el control del negocio (IFRS Foundation, 2018)."),
    ("HTTP", "Protocolo de aplicación utilizado para intercambiar solicitudes y respuestas entre clientes y servidores web (MDN Web Docs, s. f.)."),
    ("Inferencia", "Proceso de ejecutar un modelo entrenado sobre datos nuevos para producir una predicción. En el módulo de seguridad corresponde al análisis de secuencias captadas por las cámaras (PyTorch, s. f.)."),
    ("Inteligencia artificial", "Conjunto de técnicas que permite a un sistema producir resultados como predicciones, recomendaciones o decisiones a partir de entradas definidas. En este proyecto se limita a generar alertas que requieren revisión humana (OECD, 2024)."),
    ("Inventario", "Activos mantenidos para su venta en el curso ordinario del negocio. El sistema controla sus entradas, salidas, ajustes y saldos (IFRS Foundation, 2021)."),
    ("IoU", "Intersección sobre unión; métrica que compara la superposición entre una caja detectada y una caja de referencia. Se emplea para evaluar o filtrar detecciones de objetos (Ultralytics, s. f.)."),
    ("JSON", "Formato textual, liviano e independiente del lenguaje para intercambiar datos estructurados mediante objetos, arreglos y valores (Bray, 2017)."),
    ("JWT", "Formato compacto para representar un conjunto de afirmaciones o claims que puede protegerse mediante firma o cifrado (Jones et al., 2015)."),
    ("Kardex", "Registro secuencial de entradas, salidas y saldo de un producto. En el proyecto se implementa como historial digital de movimientos del inventario (IFRS Foundation, 2021)."),
    ("Machine learning", "Subcampo de la inteligencia artificial en el que los modelos aprenden patrones a partir de datos para realizar predicciones sobre casos nuevos (OECD, 2024)."),
    ("MySQL", "Sistema de gestión de bases de datos SQL que organiza información relacional en tablas y permite aplicar restricciones e integridad referencial (Oracle, 2026)."),
    ("NestJS", "Framework para aplicaciones del lado del servidor en Node.js que organiza el código mediante módulos, controladores y proveedores, con soporte para TypeScript (NestJS, s. f.-a)."),
    ("Node.js", "Entorno de ejecución de JavaScript, abierto y multiplataforma, que utiliza el motor V8 y operaciones de entrada y salida asíncronas (Node.js, s. f.)."),
    ("OpenCV", "Biblioteca de visión por computador que proporciona funciones para capturar, transformar y analizar imágenes y video (OpenCV, s. f.)."),
    ("ORM", "Capa que permite consultar y modificar una base de datos mediante modelos y operaciones del lenguaje de programación. Prisma ofrece acceso tipado, migraciones y un cliente generado (Prisma, s. f.)."),
    ("PDF", "Formato de documento portátil normalizado para representar documentos de manera independiente del software y dispositivo de origen (Adobe Systems Incorporated, 2008)."),
    ("pdfmake", "Biblioteca de JavaScript que genera documentos PDF a partir de definiciones de contenido, estilos, tablas y otros elementos (pdfmake, s. f.)."),
    ("Percepción del IGV", "Importe adicional cobrado por un agente de percepción al cliente en una venta o importación para su posterior entrega al fisco (SUNAT, s. f.-b)."),
    ("Pinia", "Biblioteca de estado para Vue.js que permite compartir datos y lógica entre componentes o páginas mediante stores (Pinia, s. f.)."),
    ("Prisma ORM", "Herramienta para Node.js y TypeScript que proporciona acceso tipado a la base de datos, modelado declarativo y migraciones (Prisma, s. f.)."),
    ("Programa de fidelización", "Conjunto de reglas que permite acumular puntos u otorgar descuentos o recompensas según las compras de un cliente (Odoo, s. f.)."),
    ("Pydantic", "Biblioteca de Python que valida y transforma datos a partir de anotaciones de tipo y modelos definidos por el desarrollador (Pydantic, s. f.)."),
    ("Python", "Lenguaje de programación de propósito general empleado en el servicio independiente de procesamiento de video e inteligencia artificial (Python Software Foundation, s. f.)."),
    ("PyTorch", "Biblioteca optimizada de tensores y aprendizaje profundo que puede ejecutarse en CPU o GPU y permite construir redes neuronales (PyTorch, s. f.)."),
    ("Red neuronal convolucional tridimensional (3D CNN)", "Modelo que aplica convoluciones sobre dimensiones espaciales y temporales, por lo que puede extraer patrones de una secuencia de video (Martínez-Mascorro et al., 2020)."),
    ("Reporte", "Presentación organizada de datos seleccionados para su consulta y análisis. El sistema generará reportes de compras, ventas, inventario, caja, deudas y clientes en PDF o XLSX (pdfmake, s. f.; ExcelJS, s. f.)."),
    ("RTSP", "Protocolo de capa de aplicación utilizado para establecer y controlar la entrega de datos con propiedades de tiempo real, como transmisiones de audio o video (Schulzrinne et al., 2016)."),
    ("Seguimiento multiobjeto", "Tarea de mantener la identidad y trayectoria de varios objetos detectados a lo largo de una secuencia de video (Zhang et al., 2022)."),
    ("Socket.IO", "Biblioteca de comunicación bidireccional basada en eventos que incluye reconexión, confirmaciones y mecanismos alternativos de transporte (Socket.IO, s. f.)."),
    ("Stock", "Cantidad disponible de un producto en un almacén o ubicación determinada. Su saldo resulta de las entradas, salidas y ajustes registrados (IFRS Foundation, 2021)."),
    ("Ticket interno", "Documento operativo generado por el sistema con el detalle y total de una venta. No constituye una boleta ni una factura y debe diferenciarse de los comprobantes de pago regulados por la SUNAT (SUNAT, s. f.-a)."),
    ("TypeScript", "Lenguaje que amplía JavaScript con comprobación estática de tipos y conserva su comportamiento de ejecución tras la compilación (Microsoft, s. f.)."),
    ("Uvicorn", "Servidor ASGI para aplicaciones Python, utilizado para ejecutar y servir la API del componente de inteligencia artificial (Uvicorn, s. f.)."),
    ("Validación de datos", "Comprobación de que la información cumple tipos, formatos y reglas antes de ser procesada. Zod y Pydantic se emplean en los límites de los servicios web (Zod, s. f.; Pydantic, s. f.)."),
    ("Venta al contado", "Operación en la que el importe se paga al momento de registrar la venta y se refleja en el medio de pago y la caja correspondiente (IFRS Foundation, 2018)."),
    ("Venta al crédito o fiado", "Operación en la que el cliente recibe los productos y queda un saldo pendiente, registrado como cuenta por cobrar hasta su cancelación mediante abonos (IFRS Foundation, 2018)."),
    ("Vite", "Herramienta de construcción para proyectos web que ofrece un servidor de desarrollo y genera recursos optimizados para producción (Vite, s. f.)."),
    ("Vue Router", "Enrutador oficial de Vue.js para construir aplicaciones de una sola página con rutas, parámetros y controles de navegación (Vue Router, s. f.)."),
    ("Vue.js", "Framework de JavaScript para crear interfaces de usuario mediante renderizado declarativo, reactividad y componentes (Vue.js, s. f.)."),
    ("Vuetify", "Framework de componentes de interfaz para Vue.js que facilita la creación de pantallas adaptables y coherentes (Vuetify, s. f.)."),
    ("WebSocket", "Protocolo que permite comunicación bidireccional entre cliente y servidor sobre una conexión persistente después de un intercambio inicial (Fette & Melnikov, 2011)."),
    ("XLSX", "Formato de libro de cálculo utilizado para intercambiar reportes tabulares con hojas, celdas, fórmulas y estilos; ExcelJS permite generarlo desde JavaScript (ExcelJS, s. f.)."),
    ("YOLOv8", "Familia de modelos de visión artificial empleada para tareas como detección, segmentación y clasificación. En el proyecto se propone para detectar personas y objetos en los fotogramas (Ultralytics, s. f.)."),
    ("Zod", "Biblioteca de validación de esquemas orientada a TypeScript que permite comprobar datos en tiempo de ejecución e inferir tipos estáticos (Zod, s. f.)."),
]


REFERENCES: list[str] = [
    "Adobe Systems Incorporated. (2008). Document management—Portable document format—Part 1: PDF 1.7. https://opensource.adobe.com/dc-acrobat-sdk-docs/pdfstandards/pdfreference1.7old.pdf",
    "Axios. (s. f.). Axios documentation. Recuperado el 28 de septiembre de 2026, de https://axios-http.com/docs/intro",
    "Bautista Coz, L. E., Blas Marcos, A. L., & Hidalgo Taipe, I. L. (2023). Sistema de punto de venta y control de inventario de la bodega J’Abdiel en la provincia de Jauja [Tesis de grado, Universidad Continental]. https://hdl.handle.net/20.500.12394/13306",
    "Bray, T. (2017). The JavaScript Object Notation (JSON) data interchange format (RFC 8259). Internet Engineering Task Force. https://doi.org/10.17487/RFC8259",
    "bwip-js. (s. f.). Barcode Writer in Pure JavaScript. GitHub. Recuperado el 28 de septiembre de 2026, de https://github.com/metafloor/bwip-js",
    "CASL. (s. f.). CASL documentation. Recuperado el 28 de septiembre de 2026, de https://casl.js.org/v6/en/",
    "ExcelJS. (s. f.). ExcelJS: Excel workbook manager. GitHub. Recuperado el 28 de septiembre de 2026, de https://github.com/exceljs/exceljs",
    "FastAPI. (s. f.). FastAPI documentation. Recuperado el 28 de septiembre de 2026, de https://fastapi.tiangolo.com/",
    "Fette, I., & Melnikov, A. (2011). The WebSocket protocol (RFC 6455). Internet Engineering Task Force. https://doi.org/10.17487/RFC6455",
    "Fielding, R. T. (2000). Architectural styles and the design of network-based software architectures [Tesis doctoral, University of California, Irvine]. https://www.ics.uci.edu/~fielding/pubs/dissertation/fielding_dissertation.pdf",
    "Frigate. (s. f.). Frigate documentation. Recuperado el 28 de septiembre de 2026, de https://docs.frigate.video/",
    "IFRS Foundation. (2018). Conceptual framework for financial reporting. https://www.ifrs.org/issued-standards/list-of-standards/conceptual-framework/",
    "IFRS Foundation. (2021). IAS 2 inventories. https://www.ifrs.org/content/dam/ifrs/publications/pdf-standards/english/2021/issued/part-a/ias-2-inventories.pdf",
    "Jones, M., Bradley, J., & Sakimura, N. (2015). JSON Web Token (JWT) (RFC 7519). Internet Engineering Task Force. https://doi.org/10.17487/RFC7519",
    "Loyverse. (s. f.). Free POS system that powers your whole business. Recuperado el 28 de septiembre de 2026, de https://loyverse.com/",
    "Martínez-Mascorro, G. A., Abreu-Pederzini, J. R., Ortiz-Bayliss, J. C., & Terashima-Marín, H. (2020). Suspicious behavior detection on shoplifting cases for crime prevention by using 3D convolutional neural networks. https://arxiv.org/abs/2005.02142",
    "MDN Web Docs. (s. f.). HTTP; CORS; Client-server overview. Recuperado el 28 de septiembre de 2026, de https://developer.mozilla.org/",
    "Microsoft. (s. f.). The TypeScript handbook. Recuperado el 28 de septiembre de 2026, de https://www.typescriptlang.org/docs/handbook/intro.html",
    "National Institute of Standards and Technology. (s. f.). Role based access control. Recuperado el 28 de septiembre de 2026, de https://csrc.nist.gov/projects/role-based-access-control",
    "NestJS. (s. f.-a). Modules, controllers and providers. Recuperado el 28 de septiembre de 2026, de https://docs.nestjs.com/",
    "NestJS. (s. f.-b). Gateways. Recuperado el 28 de septiembre de 2026, de https://docs.nestjs.com/websockets/gateways",
    "Node.js. (s. f.). Introduction to Node.js. Recuperado el 28 de septiembre de 2026, de https://nodejs.org/en/learn/getting-started/introduction-to-nodejs",
    "Odoo. (s. f.). Discount and loyalty programs. Recuperado el 28 de septiembre de 2026, de https://www.odoo.com/documentation/16.0/applications/sales/sales/products_prices/loyalty_discount.html",
    "OECD. (2024). Explanatory memorandum on the updated OECD definition of an AI system. https://doi.org/10.1787/623da898-en",
    "OpenAPI Initiative. (2021). OpenAPI specification 3.1.0. https://spec.openapis.org/oas/v3.1.0",
    "OpenCV. (s. f.). Video I/O with OpenCV overview. Recuperado el 28 de septiembre de 2026, de https://docs.opencv.org/4.x/d0/da7/videoio_overview.html",
    "Oracle. (2026). MySQL 8.4 reference manual. https://dev.mysql.com/doc/refman/8.4/en/",
    "OWASP Foundation. (s. f.-a). Authentication cheat sheet. Recuperado el 28 de septiembre de 2026, de https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html",
    "OWASP Foundation. (s. f.-b). Password storage cheat sheet. Recuperado el 28 de septiembre de 2026, de https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html",
    "Palacios Cuyatti, A. M., & Vicuña Huaylinos, C. A. (2021). Sistema web para control y gestión de ventas del Minimarket “Gandy Market” en el distrito de Comas [Tesis de grado, Universidad César Vallejo]. https://hdl.handle.net/20.500.12692/69992",
    "pdfmake. (s. f.). pdfmake documentation. Recuperado el 28 de septiembre de 2026, de https://pdfmake.github.io/docs/0.3/",
    "Pinia. (s. f.). Introduction. Recuperado el 28 de septiembre de 2026, de https://pinia.vuejs.org/introduction.html",
    "Prisma. (s. f.). Prisma ORM 7 documentation. Recuperado el 28 de septiembre de 2026, de https://www.prisma.io/docs/orm/v7",
    "Pydantic. (s. f.). Pydantic documentation. Recuperado el 28 de septiembre de 2026, de https://docs.pydantic.dev/latest/",
    "Python Software Foundation. (s. f.). The Python tutorial. Recuperado el 28 de septiembre de 2026, de https://docs.python.org/3/tutorial/",
    "PyTorch. (s. f.). PyTorch documentation. Recuperado el 28 de septiembre de 2026, de https://docs.pytorch.org/docs/stable/",
    "Schulzrinne, H., Rao, A., Lanphier, R., Westerlund, M., & Stiemerling, M. (2016). Real-Time Streaming Protocol version 2.0 (RFC 7826). Internet Engineering Task Force. https://doi.org/10.17487/RFC7826",
    "Socket.IO. (s. f.). Socket.IO documentation. Recuperado el 28 de septiembre de 2026, de https://socket.io/docs/v4/",
    "SuperAgent Project. (s. f.). SuperTest. GitHub. Recuperado el 28 de septiembre de 2026, de https://github.com/forwardemail/supertest",
    "Superintendencia Nacional de Aduanas y de Administración Tributaria. (s. f.-a). Nuevo Régimen Único Simplificado (Nuevo RUS). Recuperado el 28 de septiembre de 2026, de https://emprender.sunat.gob.pe/ruc/regimenes-tributarios-mype/nuevo-regimen-unico-simplificado-nuevo-rus",
    "Superintendencia Nacional de Aduanas y de Administración Tributaria. (s. f.-b). Comprobante de percepción. Recuperado el 28 de septiembre de 2026, de https://orientacion.sunat.gob.pe/05-comprobante-de-percepcion",
    "Ultralytics. (s. f.). Ultralytics YOLO documentation. Recuperado el 28 de septiembre de 2026, de https://docs.ultralytics.com/",
    "Uvicorn. (s. f.). Uvicorn documentation. Recuperado el 28 de septiembre de 2026, de https://www.uvicorn.org/",
    "Vite. (s. f.). Getting started. Recuperado el 28 de septiembre de 2026, de https://vite.dev/guide/",
    "Vitest. (s. f.). Vitest guide. Recuperado el 28 de septiembre de 2026, de https://vitest.dev/guide/",
    "Vue Router. (s. f.). Introduction. Recuperado el 28 de septiembre de 2026, de https://router.vuejs.org/introduction.html",
    "Vue.js. (s. f.). Introduction. Recuperado el 28 de septiembre de 2026, de https://vuejs.org/guide/introduction.html",
    "Vuetify. (s. f.). Get started with Vuetify. Recuperado el 28 de septiembre de 2026, de https://vuetifyjs.com/en/getting-started/installation/",
    "Zhang, Y., Sun, P., Jiang, Y., Yu, D., Weng, F., Yuan, Z., Luo, P., Liu, W., & Wang, X. (2022). ByteTrack: Multi-object tracking by associating every detection box. In S. Avidan et al. (Eds.), Computer Vision – ECCV 2022 (pp. 1–21). Springer. https://doi.org/10.1007/978-3-031-20047-2_1",
    "Zod. (s. f.). Zod documentation. Recuperado el 28 de septiembre de 2026, de https://zod.dev/",
]


def normalized(text: str) -> str:
    return "".join(
        char
        for char in unicodedata.normalize("NFKD", text.casefold())
        if not unicodedata.combining(char)
    )


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def configure_body(paragraph) -> None:
    paragraph.style = "APA 7MA EDICION"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    paragraph.paragraph_format.widow_control = True


def insert_heading(anchor, text: str):
    paragraph = anchor.insert_paragraph_before()
    configure_body(paragraph)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return paragraph


def insert_body(anchor, text: str):
    paragraph = anchor.insert_paragraph_before()
    configure_body(paragraph)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return paragraph


def insert_term(anchor, term: str, definition: str):
    paragraph = anchor.insert_paragraph_before()
    configure_body(paragraph)
    paragraph.paragraph_format.keep_together = True
    label = paragraph.add_run(f"{term}: ")
    label.bold = True
    label.font.name = "Times New Roman"
    label.font.size = Pt(12)
    definition_run = paragraph.add_run(definition)
    definition_run.font.name = "Times New Roman"
    definition_run.font.size = Pt(12)
    return paragraph


def configure_reference(paragraph) -> None:
    paragraph.style = "APA 7MA EDICION"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    terms = [term for term, _ in GLOSSARY]
    expected_terms = sorted(terms, key=normalized)
    if terms != expected_terms:
        mismatches = [
            (index, current, expected)
            for index, (current, expected) in enumerate(zip(terms, expected_terms))
            if current != expected
        ]
        raise ValueError(f"El glosario no está en orden alfabético: {mismatches[:5]}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    original = list(document.paragraphs)

    if len(original) < 833:
        raise ValueError("La estructura del documento no coincide con la versión esperada.")
    if "Marco" not in original[189].text or "CAPÍTULO III" not in original[624].text:
        raise ValueError("No se localizaron los límites esperados de la sección 2.4.")

    marco_heading = original[189]
    fundamento_heading = original[190]
    conceptos_heading = original[328]
    chapter_three = original[624]
    references_title = original[829]
    references_placeholder = original[830]
    references_section_break = original[831]
    annexes_title = original[832]

    marco_heading.text = "Marco teórico y conceptual."
    fundamento_heading.text = "Fundamento teórico del proyecto de mejora."
    conceptos_heading.text = "Conceptos y términos utilizados."
    for heading in (marco_heading, fundamento_heading, conceptos_heading):
        heading.paragraph_format.line_spacing = 2
        heading.paragraph_format.space_before = Pt(0)
        heading.paragraph_format.space_after = Pt(0)
        heading.paragraph_format.keep_with_next = True
        for run in heading.runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    for paragraph in original[191:328]:
        remove_paragraph(paragraph)
    for paragraph in original[329:624]:
        remove_paragraph(paragraph)

    insert_body(
        conceptos_heading,
        "El fundamento teórico articula los procesos comerciales de la tienda con las tecnologías empleadas en Innova Store y con el módulo de seguridad propuesto. Las siguientes bases explican cómo cada componente contribuye al registro, la trazabilidad, la comunicación y el análisis de la información.",
    )
    for heading, paragraphs in THEORY_SECTIONS:
        insert_heading(conceptos_heading, heading)
        for text in paragraphs:
            insert_body(conceptos_heading, text)

    insert_body(
        chapter_three,
        "Los términos se presentan en orden alfabético y cada definición incluye la fuente que sustenta su uso técnico u operativo dentro del proyecto.",
    )
    for term, definition in GLOSSARY:
        insert_term(chapter_three, term, definition)

    chapter_three.paragraph_format.page_break_before = True

    references_title.text = "REFERENCIAS BIBLIOGRÁFICAS"
    references_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    references_title.paragraph_format.line_spacing = 2
    references_title.paragraph_format.space_before = Pt(0)
    references_title.paragraph_format.space_after = Pt(0)
    references_title.paragraph_format.keep_with_next = True
    for run in references_title.runs:
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    references_placeholder.text = REFERENCES[0]
    configure_reference(references_placeholder)
    for reference in REFERENCES[1:]:
        paragraph = references_section_break.insert_paragraph_before(reference)
        configure_reference(paragraph)
    annexes_title.paragraph_format.page_break_before = False

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.core_properties.title = (
        "Marco teórico y conceptual del sistema de gestión para Minimarket Tienda Perez"
    )
    document.core_properties.subject = (
        "Actualización de la sección 2.4 con tecnologías, conceptos y referencias"
    )
    document.save(OUTPUT)

    check = Document(OUTPUT)
    texts = [paragraph.text for paragraph in check.paragraphs]
    required = [
        "Frontend con Vue.js y TypeScript",
        "Comunicación en tiempo real mediante WebSocket y Socket.IO",
        "Modelo propuesto: YOLOv8, ByteTrack y red 3D CNN",
        "Ticket interno:",
        "REFERENCIAS BIBLIOGRÁFICAS",
    ]
    for value in required:
        if not any(value in text for text in texts):
            raise ValueError(f"Falta contenido requerido: {value}")
    if any("gestión de cocheras" in text.casefold() for text in texts[189:]):
        raise ValueError("Permaneció contenido del marco teórico de cocheras.")

    print(f"Documento guardado: {OUTPUT}")
    print(f"Párrafos: {len(check.paragraphs)}")
    print(f"Términos en 2.4.2: {len(GLOSSARY)}")
    print(f"Referencias: {len(REFERENCES)}")
    print(f"Tamaño: {os.path.getsize(OUTPUT)} bytes")


if __name__ == "__main__":
    main()
