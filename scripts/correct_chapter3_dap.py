from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from update_thesis_chapter4 import load_fonts, set_run_font


SOURCE = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\output\documentos"
    r"\Tesis_JosephKleynMamaniPerez_capitulos3y4_actualizado.docx"
)
OUTPUT = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\output\documentos"
    r"\Tesis_JosephKleynMamaniPerez_capitulos3y4_diagramas_corregidos.docx"
)
ASSET_DIR = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\tmp\chapters3-4-assets"
)
DAP_IMAGE = ASSET_DIR / "chapter3-dap-actual.png"


ACTIVITIES = [
    ("1", "Revisar existencias en estantes", "O", "Control basado en observación y memoria"),
    ("2", "Preparar el pedido al proveedor", "O", "La cantidad se determina manualmente"),
    ("3", "Trasladar la mercadería hasta la tienda", "T", "Transporte realizado por el proveedor"),
    ("4", "Verificar cantidades y estado de los productos", "I", "Conteo y revisión visual"),
    ("5", "Registrar mentalmente o anotar la compra", "O", "No se actualiza el inventario"),
    ("6", "Revisar el comprobante y la percepción", "I", "Validación manual de importes"),
    ("7", "Esperar la conciliación o el registro posterior", "D", "La información puede quedar pendiente"),
    ("8", "Archivar los documentos del proveedor", "A", "Archivador físico"),
    ("9", "Trasladar los productos a estantes o refrigeradoras", "T", "No genera un movimiento digital"),
    ("10", "Seleccionar productos y calcular la venta", "O", "Suma manual"),
    ("11", "Verificar el monto total", "I", "Revisión antes del cobro"),
    ("12", "Elaborar el ticket interno", "O", "Papel y lapicero"),
    ("13", "Cobrar o registrar la venta al fiado", "O", "El fiado se anota en un cuaderno"),
    ("14", "Guardar el ticket o la anotación del fiado", "A", "Registros físicos separados"),
]


def text_center(draw: ImageDraw.ImageDraw, xy, text: str, font, fill="#111111") -> None:
    draw.text(xy, text, font=font, fill=fill, anchor="mm", align="center")


def wrap(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_dap_symbol(draw: ImageDraw.ImageDraw, kind: str, center: tuple[int, int], size: int = 48) -> None:
    x, y = center
    half = size // 2
    color = "#17365D"
    fill = "#DDEBF7"
    if kind == "O":
        draw.ellipse((x - half, y - half, x + half, y + half), fill=fill, outline=color, width=4)
    elif kind == "I":
        draw.rectangle((x - half, y - half, x + half, y + half), fill="#FFF2CC", outline=color, width=4)
    elif kind == "T":
        points = [
            (x - half, y - 11), (x + 5, y - 11), (x + 5, y - half),
            (x + half, y), (x + 5, y + half), (x + 5, y + 11), (x - half, y + 11),
        ]
        draw.polygon(points, fill="#D9EAD3", outline=color)
        draw.line(points + [points[0]], fill=color, width=4)
    elif kind == "D":
        draw.line((x - half, y - half, x - half, y + half), fill=color, width=4)
        draw.line((x - half, y - half, x, y - half), fill=color, width=4)
        draw.line((x - half, y + half, x, y + half), fill=color, width=4)
        draw.arc((x - half, y - half, x + half, y + half), -90, 90, fill=color, width=4)
        draw.pieslice((x - half + 4, y - half + 4, x + half - 4, y + half - 4), -90, 90, fill="#FCE4D6")
    elif kind == "A":
        draw.polygon([(x - half, y - half), (x + half, y - half), (x, y + half)], fill="#E2F0D9", outline=color)
        draw.line([(x - half, y - half), (x + half, y - half), (x, y + half), (x - half, y - half)], fill=color, width=4)


def create_dap_image(path: Path) -> None:
    width, height = 2200, 1450
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = load_fonts(42, bold=True)
    header_font = load_fonts(28, bold=True)
    body_font = load_fonts(27)
    small_font = load_fonts(24)
    small_bold = load_fonts(24, bold=True)
    line = "#707070"
    blue = "#1F4E78"

    text_center(draw, (width // 2, 50), "DIAGRAMA DE ANÁLISIS DEL PROCESO ACTUAL (DAP)", title_font, "#17365D")
    draw.rectangle((40, 90, 2160, 245), outline=line, width=3)
    draw.line((40, 142, 2160, 142), fill=line, width=2)
    draw.line((40, 194, 2160, 194), fill=line, width=2)
    draw.text((60, 103), "Empresa: Minimarket Tienda Perez", font=body_font, fill="#111111")
    draw.text((60, 155), "Proceso: compras, percepciones, ventas y control", font=body_font, fill="#111111")
    draw.text((60, 207), "Método: actual", font=body_font, fill="#111111")

    counts = {"O": 6, "I": 3, "T": 2, "D": 1, "A": 2}
    labels = {"O": "Operación", "I": "Inspección", "T": "Transporte", "D": "Demora", "A": "Almacén"}
    for index, kind in enumerate(("O", "I", "T", "D", "A")):
        x = 1215 + index * 185
        draw_dap_symbol(draw, kind, (x, 127), 38)
        draw.text((x + 35, 107), f"{labels[kind]}: {counts[kind]}", font=small_font, fill="#111111")

    table_top = 275
    header_height = 95
    row_height = 72
    xs = [40, 120, 1010, 1120, 1230, 1340, 1450, 1560, 2160]
    draw.rectangle((40, table_top, 2160, table_top + header_height), fill=blue, outline=line, width=3)
    for x in xs:
        draw.line((x, table_top, x, table_top + header_height + row_height * len(ACTIVITIES)), fill=line, width=2)
    text_center(draw, ((xs[0] + xs[1]) // 2, table_top + header_height // 2), "N.º", header_font, "#FFFFFF")
    text_center(draw, ((xs[1] + xs[2]) // 2, table_top + header_height // 2), "Descripción", header_font, "#FFFFFF")
    for i, kind in enumerate(("O", "I", "T", "D", "A"), start=2):
        draw_dap_symbol(draw, kind, ((xs[i] + xs[i + 1]) // 2, table_top + 36), 34)
        text_center(draw, ((xs[i] + xs[i + 1]) // 2, table_top + 77), kind, small_bold, "#FFFFFF")
    text_center(draw, ((xs[7] + xs[8]) // 2, table_top + header_height // 2), "Observación", header_font, "#FFFFFF")

    for row_index, (number, description, kind, observation) in enumerate(ACTIVITIES):
        y1 = table_top + header_height + row_index * row_height
        y2 = y1 + row_height
        if row_index % 2:
            draw.rectangle((40, y1, 2160, y2), fill="#F2F7FB")
        draw.line((40, y2, 2160, y2), fill=line, width=2)
        text_center(draw, ((xs[0] + xs[1]) // 2, (y1 + y2) // 2), number, body_font)
        desc_lines = wrap(draw, description, body_font, xs[2] - xs[1] - 28)
        draw.multiline_text((xs[1] + 14, (y1 + y2) // 2), "\n".join(desc_lines), font=body_font, fill="#111111", anchor="lm", spacing=3)
        kind_index = {"O": 2, "I": 3, "T": 4, "D": 5, "A": 6}[kind]
        draw_dap_symbol(draw, kind, ((xs[kind_index] + xs[kind_index + 1]) // 2, (y1 + y2) // 2), 42)
        obs_lines = wrap(draw, observation, small_font, xs[8] - xs[7] - 24)
        draw.multiline_text((xs[7] + 12, (y1 + y2) // 2), "\n".join(obs_lines), font=small_font, fill="#111111", anchor="lm", spacing=2)

    footer_y = table_top + header_height + row_height * len(ACTIVITIES) + 35
    draw.text((45, footer_y), "Resumen: 6 operaciones, 3 inspecciones, 2 transportes, 1 demora y 2 almacenamientos.", font=small_bold, fill="#17365D")
    draw.text((45, footer_y + 38), "El tiempo y la distancia no se consignan porque el negocio aún no dispone de mediciones históricas verificables.", font=small_font, fill="#404040")
    image.save(path)


def set_paragraph_text(paragraph, text: str, *, italic: bool = False) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    set_run_font(run, italic=italic)


def replace_figure(document: Document) -> None:
    paragraphs = document.paragraphs
    figure_number = next(
        p for p in paragraphs
        if p.text.strip() == "Figura 2" and document._element.body.index(p._p) > 100
    )
    index = paragraphs.index(figure_number)
    title = paragraphs[index + 1]
    image_paragraph = paragraphs[index + 2]
    note = paragraphs[index + 3]

    set_paragraph_text(title, "Diagrama de Análisis del Proceso actual de compras, percepciones, ventas y control", italic=True)
    for child in list(image_paragraph._p):
        if child.tag != qn("w:pPr"):
            image_paragraph._p.remove(child)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_together = True
    image_paragraph.add_run().add_picture(str(DAP_IMAGE), width=Inches(6.15))

    for child in list(note._p):
        if child.tag != qn("w:pPr"):
            note._p.remove(child)
    lead = note.add_run("Nota. ")
    set_run_font(lead, italic=True)
    body = note.add_run("Elaboración propia con base en la simbología DAP descrita por SENATI (2022).")
    set_run_font(body)

    for p in document.paragraphs:
        if p.text.startswith("En la venta, los productos se seleccionan"):
            text = (
                "En la venta, los productos se seleccionan y suman manualmente. El ticket interno se escribe con papel y lapicero, "
                "y cuando la operación es al crédito el cliente, el monto y los abonos se anotan en un cuaderno. Los clientes frecuentes "
                "se reconocen de memoria y las cámaras se consultan por separado. La Figura 2 presenta el DAP actual mediante símbolos "
                "para operación, inspección, transporte, demora y almacenamiento (SENATI, 2022)."
            )
            set_paragraph_text(p, text)
            break

    for p in document.paragraphs:
        if p.text.strip() == "Secuencia y riesgos del proceso manual actual":
            set_paragraph_text(p, "Detalle y riesgos de las actividades del DAP actual", italic=True)
            break


def add_reference(document: Document) -> None:
    if any("Mejora de métodos en el trabajo" in p.text for p in document.paragraphs):
        return
    target = next(p for p in document.paragraphs if p.text.startswith("Socket.IO."))
    p = target.insert_paragraph_before()
    p.style = target.style
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(
        "Servicio Nacional de Adiestramiento en Trabajo Industrial. (2022). "
        "Mejora de métodos en el trabajo: Manual de aprendizaje profesional técnico (CGEU-241). SENATI."
    )
    set_run_font(run)


def prevent_table_row_splitting(document: Document) -> None:
    for table in document.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    create_dap_image(DAP_IMAGE)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    replace_figure(document)
    add_reference(document)
    prevent_table_row_splitting(document)
    document.core_properties.subject = (
        "Capítulos III y IV actualizados con Diagrama de Análisis del Proceso según simbología SENATI"
    )
    document.save(OUTPUT)
    print(f"Documento corregido: {OUTPUT}")
    print(f"Párrafos: {len(document.paragraphs)} | Tablas: {len(document.tables)} | Imágenes: {len(document.inline_shapes)}")


if __name__ == "__main__":
    main()
