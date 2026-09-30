from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\output\documentos"
    r"\Tesis_JosephKleynMamaniPerez_capitulos3y4_diagramas_corregidos.docx"
)

PAGES = [
    [
        ("CAPÍTULO I", "8", 0),
        ("1.1 Razón social", "8", 1),
        ("1.2 Misión, visión, objetivos y valores de la empresa", "9", 1),
        ("1.3 Productos, mercado y clientes", "10", 1),
        ("1.4 Estructura de la organización", "11", 1),
        ("1.5 Otra información relevante de la empresa donde se desarrolla el proyecto", "12", 1),
        ("CAPÍTULO II", "13", 0),
        ("2.1 Identificación del problema técnico en la empresa", "13", 1),
        ("2.2 Objetivos del proyecto de mejora", "16", 1),
        ("2.3 Antecedentes del proyecto de mejora", "17", 1),
        ("Justificación del proyecto de mejora", "20", 1),
        ("2.4 Marco teórico y conceptual", "22", 1),
        ("2.4.1 Fundamento teórico del proyecto de mejora", "22", 2),
        ("2.4.2 Conceptos y términos utilizados", "31", 2),
    ],
    [
        ("CAPÍTULO III", "39", 0),
        ("3.1 Diagrama del proceso, mapa del flujo de valor y/o diagrama de operación actual", "39", 1),
        ("3.2 Efectos del problema en el área de trabajo o en los resultados de la empresa", "41", 1),
        ("3.3 Análisis de las causas raíz que generan el problema", "42", 1),
        ("3.4 Priorización de causas raíz", "44", 1),
        ("CAPÍTULO IV", "46", 0),
        ("4.1 Plan de acción de la mejora propuesta", "46", 1),
        ("4.2 Consideraciones técnicas, operativas y ambientales para la implementación de la mejora", "47", 1),
        ("4.3 Recursos técnicos para implementar la mejora propuesta", "48", 1),
        ("4.4 Diagrama del proceso, mapa del flujo de valor y diagrama de operación de la situación mejorada", "50", 1),
        ("4.5 Control de compras y percepciones", "55", 1),
        ("4.6 Aspectos limitantes para la implementación de la mejora", "56", 1),
    ],
    [
        ("CAPÍTULO V", "57", 0),
        ("5.1 Costo de materiales", "57", 1),
        ("5.2 Costo de mano de obra", "57", 1),
        ("5.3 Costo de máquinas, herramientas y equipos", "57", 1),
        ("5.4 Otros costos de implementación", "57", 1),
        ("5.5 Costo total de la implementación de la mejora", "57", 1),
        ("CAPÍTULO VI", "58", 0),
        ("6.1 Beneficio técnico y/o económico esperado", "58", 1),
        ("6.2 Relación beneficio/costo", "58", 1),
        ("CAPÍTULO VII", "59", 0),
        ("7.1 Conclusiones respecto a los objetivos del proyecto de mejora", "59", 1),
        ("CAPÍTULO VIII", "60", 0),
        ("8.1 Recomendaciones para la empresa respecto del proyecto de mejora", "60", 1),
        ("REFERENCIAS BIBLIOGRÁFICAS", "61", 0),
        ("ANEXOS", "66", 0),
    ],
]


def format_entry(paragraph, title: str, page: str, level: int, page_break_before: bool) -> None:
    paragraph.style = "Normal"
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.18 * level)
    fmt.first_line_indent = Inches(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.0
    fmt.keep_together = True
    fmt.page_break_before = page_break_before
    fmt.tab_stops.add_tab_stop(Inches(6.05 - 0.18 * level), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = paragraph.add_run(f"{title}\t{page}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), "Times New Roman")


def main() -> None:
    document = Document(DOCX)
    body = document._element.body
    index_heading = next(p for p in document.paragraphs if p.text.strip().casefold() == "índice")
    chapter_one = next(p for p in document.paragraphs if p.text.strip().casefold() in {"capitulo i", "capítulo i"})
    children = list(body)
    heading_index = children.index(index_heading._p)
    chapter_index = children.index(chapter_one._p)

    section_break = None
    for element in reversed(children[heading_index + 1:chapter_index]):
        if element.tag == qn("w:p"):
            ppr = element.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                section_break = element
                break
    if section_break is None:
        raise RuntimeError("No se encontró el salto de sección posterior al índice.")

    children = list(body)
    break_index = children.index(section_break)
    for element in children[heading_index + 1:break_index]:
        body.remove(element)

    created = []
    for page_index, entries in enumerate(PAGES):
        for entry_index, (title, page, level) in enumerate(entries):
            paragraph = document.add_paragraph()
            format_entry(paragraph, title, page, level, entry_index == 0 and page_index > 0)
            created.append(paragraph._p)

    insertion_index = list(body).index(section_break)
    for offset, element in enumerate(created):
        body.insert(insertion_index + offset, element)

    document.save(DOCX)
    print(f"Índice actualizado: {DOCX}")


if __name__ == "__main__":
    main()
