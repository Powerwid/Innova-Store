from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from update_thesis_chapter4 import (
    create_ai_diagram,
    create_architecture_diagram,
    create_process_diagram,
    format_body,
    insert_body,
    insert_caption,
    insert_figure,
    insert_note,
    insert_subheading,
    insert_table,
    load_fonts,
    rounded_box,
    arrow,
    set_run_font,
)


SOURCE = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\output\documentos"
    r"\Tesis_JosephKleynMamaniPerez_capitulo4_actualizado.docx"
)
OUTPUT = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\output\documentos"
    r"\Tesis_JosephKleynMamaniPerez_capitulos3y4_actualizado.docx"
)
ER_SOURCE = Path(r"C:\Users\Power\Downloads\diagrama_store_w.png")
ASSET_DIR = Path(
    r"C:\Users\Power\Documents\Proyecto-Empresa\Innova-Store\tmp\chapters3-4-assets"
)


CURRENT_PROCESS_ROWS = [
    ["1", "Revisar la necesidad de reposición", "Observación de estantes o memoria del propietario", "No existe un stock actualizado ni un punto de reposición verificable."],
    ["2", "Realizar el pedido al proveedor", "Llamada, visita o nota informal", "La cantidad solicitada puede no coincidir con la demanda real."],
    ["3", "Recibir mercadería", "Conteo visual y comprobante del proveedor", "Pueden omitirse diferencias, productos dañados o unidades no entregadas."],
    ["4", "Revisar el comprobante recibido", "Factura, boleta u otro documento físico", "La información queda dispersa y resulta difícil localizarla por proveedor o fecha."],
    ["5", "Identificar la percepción", "Revisión manual del comprobante de percepción", "Existe riesgo de no registrar base, porcentaje, monto, periodo o saldo aplicable."],
    ["6", "Calcular el costo de la compra", "Calculadora o anotación manual", "El costo puede excluir la percepción u otros importes relacionados."],
    ["7", "Guardar documentos", "Archivador físico", "La búsqueda y conciliación posterior requieren tiempo y pueden perderse documentos."],
    ["8", "Colocar productos para la venta", "Distribución en estantes y refrigeradoras", "El inventario no se actualiza cuando ingresa la mercadería."],
    ["9", "Atender la venta", "Selección y suma manual", "Puede existir error en cantidades, precios o monto total."],
    ["10", "Elaborar ticket", "Papel y lapicero", "La escritura manual demora y dificulta conservar un historial de ventas."],
    ["11", "Registrar un fiado", "Cuaderno", "Los abonos y saldos pueden quedar incompletos o resultar difíciles de comprobar."],
    ["12", "Controlar caja y gastos", "Cuadernos, comprobantes y memoria", "No se obtiene un saldo consolidado ni trazabilidad de cada movimiento."],
    ["13", "Reconocer clientes frecuentes", "Memoria de los propietarios", "Las promociones no se sustentan en un historial objetivo de compras."],
    ["14", "Revisar resultados", "Conteo y revisión manual", "No existen reportes oportunos de compras, percepciones, ventas, stock y deudas."],
]

EFFECT_ROWS = [
    ["Compras y percepciones", "Documentos físicos y cálculos aislados", "Omisión de percepciones, dificultad para conocer el monto acumulado y falta de respaldo ordenado por periodo.", "Alto"],
    ["Inventario", "Entradas y salidas sin actualización inmediata", "Diferencias entre la existencia real y la estimada, compras innecesarias o quiebres de stock.", "Alto"],
    ["Ventas y tickets", "Suma y escritura manual", "Cobros incorrectos, demora en la atención y ausencia de historial detallado.", "Alto"],
    ["Caja y gastos", "Registros distribuidos", "Dificultad para conciliar ingresos, egresos, pagos a proveedores y efectivo disponible.", "Alto"],
    ["Fiados", "Deudas y abonos anotados en cuaderno", "Saldos desactualizados, cobros tardíos y desacuerdos con clientes.", "Alto"],
    ["Clientes frecuentes", "Reconocimiento basado en la memoria", "Promociones poco verificables y pérdida de oportunidades de fidelización.", "Medio"],
    ["Seguridad", "Cámaras consultadas de manera independiente", "Respuesta tardía ante aglomeraciones o posibles acciones sospechosas.", "Medio"],
    ["Toma de decisiones", "Información incompleta y no consolidada", "No se dispone de reportes confiables para planificar compras, precios y promociones.", "Alto"],
]

WHY_ROWS = [
    ["1", "¿Por qué no se conoce con precisión el monto de las compras y percepciones?", "Porque los comprobantes y montos se revisan y archivan manualmente."],
    ["2", "¿Por qué el registro manual no permite un control oportuno?", "Porque no existe un historial centralizado que relacione proveedor, compra, comprobante, percepción, pago e inventario."],
    ["3", "¿Por qué la información no está relacionada?", "Porque cada operación se anota en cuadernos, documentos físicos o registros separados."],
    ["4", "¿Por qué se mantienen registros separados?", "Porque el negocio no dispone de una herramienta adaptada a su forma de trabajo y a la facilidad de uso requerida por los propietarios."],
    ["5", "¿Por qué no se implementó antes una herramienta adecuada?", "Porque no se habían formalizado los procesos, datos, controles y prioridades necesarios para desarrollar una solución integral."],
]

CRITERIA_ROWS = [
    ["Frecuencia", "Cuántas veces se presenta la causa durante la operación cotidiana.", "1 = muy baja; 5 = muy alta"],
    ["Impacto", "Consecuencia económica, operativa o administrativa que puede ocasionar.", "1 = menor; 5 = crítica"],
    ["Urgencia", "Necesidad de intervenir para evitar que el problema continúe o aumente.", "1 = puede esperar; 5 = inmediata"],
]

PRIORITY_ROWS = [
    ["Ausencia de un sistema integral", "5", "5", "5", "15", "Muy alta"],
    ["Control manual de compras y percepciones", "5", "5", "5", "15", "Muy alta"],
    ["Registros manuales no estandarizados", "5", "5", "4", "14", "Muy alta"],
    ["Stock sin actualización inmediata", "4", "5", "4", "13", "Alta"],
    ["Tickets y cálculos manuales", "4", "4", "4", "12", "Alta"],
    ["Fiados registrados en cuaderno", "4", "4", "3", "11", "Alta"],
    ["Falta de historial de clientes frecuentes", "3", "3", "3", "9", "Media"],
    ["Cámaras no integradas al sistema", "3", "3", "2", "8", "Media"],
]

PLAN_ROWS = [
    ["1", "Confirmar alcance y requisitos", "Validar compras, percepciones, inventario, ventas, caja, fiados, clientes y seguridad.", "Responsable del proyecto y propietarios", "Requisitos aprobados"],
    ["2", "Diseñar la solución", "Definir arquitectura, modelo de datos, permisos, interfaces y criterios de aceptación.", "Responsable del proyecto", "Diseño técnico"],
    ["3", "Configurar la base del sistema", "Preparar autenticación, usuarios, roles, permisos, sucursal y parámetros.", "Responsable del proyecto", "Acceso seguro"],
    ["4", "Implementar compras y percepciones", "Registrar proveedor, comprobante, base, porcentaje, percepción, pago y acumulado mensual.", "Responsable del proyecto", "Compras y percepciones trazables"],
    ["5", "Implementar inventario", "Registrar recepción, movimientos, stock mínimo y ajustes con evidencia.", "Responsable del proyecto", "Kardex actualizado"],
    ["6", "Implementar ventas y caja", "Calcular importes, registrar medios de pago, generar tickets internos y actualizar caja.", "Responsable del proyecto", "Venta trazable"],
    ["7", "Implementar fiados y clientes", "Controlar deudas, abonos, saldos, compras frecuentes, promociones y recompensas.", "Responsable del proyecto y propietarios", "Cuentas por cobrar e historial"],
    ["8", "Incorporar reportes y alertas", "Generar PDF, XLSX, códigos, alertas de stock y avisos configurables sobre compras acumuladas.", "Responsable del proyecto", "Información para decisión"],
    ["9", "Incorporar eventos en tiempo real", "Sincronizar ventas, caja, inventario, deudas y alertas mediante Socket.IO.", "Responsable del proyecto", "Pantallas actualizadas"],
    ["10", "Integrar cámaras e IA", "Conectar flujos autorizados y emitir alertas que siempre requieran revisión humana.", "Responsable del proyecto y propietarios", "Seguridad integrada"],
    ["11", "Probar, migrar y desplegar", "Ejecutar pruebas, cargar datos, capacitar, respaldar y poner en operación.", "Responsable del proyecto y usuarios", "Sistema validado"],
]

CONSIDERATION_ROWS = [
    ["Arquitectura", "Separar interfaz, API, base de datos y servicio de video.", "Mantener módulos independientes y contratos de comunicación definidos."],
    ["Seguridad", "Proteger credenciales, información comercial y datos de clientes.", "Aplicar autenticación, roles, permisos, hash de contraseñas y auditoría."],
    ["Integridad", "Evitar operaciones aplicadas parcialmente.", "Usar validaciones y transacciones para actualizar compra, inventario, venta, caja y deuda."],
    ["Percepciones y Nuevo RUS", "Controlar documentos recibidos, importe percibido y compras acumuladas.", "Registrar base, porcentaje, monto, periodo y saldo; emitir alertas configurables y reportes para revisión del propietario o contador."],
    ["Disponibilidad", "Una falla del servicio de cámaras no debe detener la gestión comercial.", "Desacoplar el servicio de video y mostrar su estado sin bloquear compras o ventas."],
    ["Usabilidad", "Los propietarios son adultos mayores.", "Emplear textos claros, controles visibles, pocos pasos y confirmaciones."],
    ["Operación", "Los datos iniciales provienen de cuadernos y comprobantes físicos.", "Depurar catálogos, saldos y documentos antes de la carga inicial."],
    ["Ambiental", "La solución utiliza energía y documentos impresos.", "Digitalizar registros e imprimir tickets internos solo cuando sea necesario."],
    ["Videovigilancia", "El análisis de imágenes puede generar falsas alertas.", "Limitar accesos, ajustar umbrales y exigir revisión humana."],
]

SOFTWARE_ROWS = [
    ["Interfaz web", "Vue.js, TypeScript, Vuetify, Pinia, Vue Router y Vite", "Pantallas, navegación, estado compartido y compilación del frontend"],
    ["Servidor", "Node.js, NestJS y TypeScript", "API, reglas del negocio, autenticación y módulos"],
    ["Datos", "MySQL y Prisma ORM", "Persistencia relacional, consultas, migraciones y transacciones"],
    ["Comunicación", "HTTP, JSON, Axios, WebSocket y Socket.IO", "Intercambio de datos y eventos en tiempo real"],
    ["Documentos", "pdfmake, ExcelJS y bwip-js", "Tickets internos, reportes PDF/XLSX y códigos de barras"],
    ["Inteligencia artificial", "Python, FastAPI, OpenCV, PyTorch, YOLOv8, ByteTrack y 3D CNN", "Captura, detección, seguimiento y clasificación de video"],
    ["Calidad", "Vitest, Supertest, Oxlint y Prettier", "Pruebas, validación de endpoints y consistencia del código"],
]

HARDWARE_ROWS = [
    ["Equipo de atención", "Computadora existente o equivalente y navegador actualizado", "Compras, ventas, caja y consultas"],
    ["Servidor", "Equipo local o alojamiento según el despliegue", "Backend y base de datos"],
    ["Impresora de tickets", "Impresora térmica compatible", "Entrega opcional del ticket interno"],
    ["Lector de códigos", "Lector USB o cámara compatible", "Búsqueda rápida de productos"],
    ["Cámaras", "Cámaras IP existentes con acceso autorizado", "Visualización y análisis de secuencias"],
    ["Red", "Router, cableado y conexión estable", "Comunicación entre equipos, cámaras y servidor"],
    ["Respaldo eléctrico", "UPS o estabilizador recomendado", "Continuidad y protección de datos"],
    ["Procesamiento de IA", "CPU compatible; GPU opcional", "Inferencia y procesamiento de video"],
]

HUMAN_ROWS = [
    ["Responsable del proyecto", "Análisis, diseño, programación, pruebas, despliegue y documentación"],
    ["Propietarios", "Validación de reglas, datos iniciales, promociones, límites y resultados"],
    ["Usuarios de atención", "Pruebas de ventas, caja, fiados y facilidad de uso"],
    ["Asesor especializado", "Revisión técnica, contable o de seguridad cuando corresponda"],
]

IMPROVED_PROCESS_ROWS = [
    ["1", "Iniciar sesión", "Autenticación, rol, permiso y sucursal", "Sesión autorizada"],
    ["2", "Registrar proveedor y compra", "Datos obligatorios, duplicados y totales", "Compra identificada"],
    ["3", "Registrar comprobante y percepción", "Tipo, serie, número, fecha, base, porcentaje, monto y periodo", "Documento trazable"],
    ["4", "Validar recepción", "Cantidades, costos, diferencias y estado", "Mercadería aceptada"],
    ["5", "Actualizar inventario", "Transacción de entrada y kardex", "Stock actualizado"],
    ["6", "Actualizar acumulado mensual", "Compras del periodo y umbrales configurados", "Alerta o estado del periodo"],
    ["7", "Seleccionar productos para venta", "Código, precio, unidad y existencia", "Detalle válido"],
    ["8", "Calcular y confirmar total", "Subtotales, descuentos y monto final", "Importe correcto"],
    ["9", "Registrar contado o fiado", "Medio de pago, caja o cuenta por cobrar", "Operación económica"],
    ["10", "Generar ticket interno", "Correlativo, detalle, cantidades y total", "Constancia operativa"],
    ["11", "Actualizar stock, caja, deuda y cliente", "Transacción y auditoría", "Saldos consistentes"],
    ["12", "Consultar reportes y alertas", "Filtros, permisos y revisión humana", "Información de supervisión"],
]

PERCEPTION_ROWS = [
    ["Proveedor", "RUC o documento, razón social y estado obligatorios; control de duplicados.", "Proveedor asociado a la compra.", "Identificar el origen de la mercadería y de la percepción."],
    ["Documento de compra", "Tipo, serie, número, fecha, subtotal, IGV y total; combinación única por proveedor.", "Comprobante emitido por el proveedor.", "Conservar trazabilidad documental."],
    ["Comprobante de percepción", "Número, fecha y vínculo obligatorio con la compra correspondiente.", "Registro de percepción relacionado.", "Evitar documentos sin sustento o duplicados."],
    ["Base y porcentaje", "Valores numéricos no negativos; porcentaje dentro de un rango configurable.", "Base imponible y tasa aplicada.", "Reproducir y verificar el cálculo."],
    ["Monto percibido", "Cálculo automático y posibilidad de ajuste justificado según documento.", "Importe de percepción.", "Conocer el monto realmente pagado por este concepto."],
    ["Total de compra", "Conciliación entre subtotal, IGV, percepción y total cancelado.", "Costo total asociado.", "Evitar diferencias entre documento, pago y egreso."],
    ["Pago o egreso", "Medio de pago, fecha, caja y usuario obligatorios cuando corresponda.", "Movimiento económico vinculado.", "Conciliar compras con la caja."],
    ["Acumulado mensual", "Suma automática de adquisiciones por periodo y sucursal.", "Indicador mensual de compras.", "Dar seguimiento preventivo al límite aplicable."],
    ["Alertas de categoría y límite", "Avisos configurables al aproximarse o superar S/ 5 000 y S/ 8 000 mensuales.", "Notificación y registro de revisión.", "Apoyar el control del Nuevo RUS sin sustituir la evaluación tributaria."],
    ["Reporte del periodo", "Filtro por fechas, proveedor y estado; exportación PDF y XLSX.", "Resumen de compras, percepciones y saldo informado.", "Facilitar revisión, compensación o solicitud de devolución cuando corresponda."],
]

RISK_ROWS = [
    ["Información inicial incompleta o duplicada", "Media", "Alta", "Depurar productos, proveedores, stock, clientes y deudas.", "Aplicar ajustes trazables y conservar respaldo."],
    ["Comprobantes o percepciones mal digitados", "Media", "Alta", "Validaciones, unicidad, cálculo automático y revisión contra el documento.", "Corregir mediante una operación auditada."],
    ["Configuración tributaria desactualizada", "Media", "Alta", "Mantener parámetros editables y revisión periódica con información oficial o asesoría contable.", "Actualizar umbrales y volver a evaluar los periodos afectados."],
    ["Interrupción eléctrica o de red", "Media", "Alta", "Respaldos, UPS y monitoreo.", "Aplicar registro temporal y regularizar las operaciones."],
    ["Cámaras incompatibles o inestables", "Media", "Media", "Probar RTSP, resolución, credenciales y ancho de banda.", "Integrar solo dispositivos compatibles."],
    ["Rendimiento insuficiente para video", "Media", "Alta", "Ajustar resolución, FPS, zonas y transmisiones.", "Procesar por intervalos o incorporar hardware."],
    ["Falsas alertas del modelo", "Alta", "Alta", "Medir resultados, ajustar umbrales y entrenar con datos representativos.", "Exigir siempre revisión humana."],
    ["Exposición de datos o imágenes", "Baja", "Alta", "Roles, cifrado, credenciales protegidas y conservación limitada.", "Revocar accesos y revisar auditoría."],
    ["Dificultad de adaptación", "Media", "Alta", "Interfaz simple y capacitación con tareas reales.", "Despliegue gradual y acompañamiento."],
    ["Confusión del ticket interno", "Media", "Alta", "Indicar que no es boleta ni factura.", "Corregir formato y reforzar capacitación."],
]


def section_heading(anchor: Paragraph, text: str) -> Paragraph:
    p = anchor.insert_paragraph_before()
    p.style = "List Paragraph"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, bold=True)
    return p


def chapter_subtitle(anchor: Paragraph, text: str) -> Paragraph:
    p = anchor.insert_paragraph_before()
    p.style = "Heading 4"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        set_run_font(r, bold=True)
    run = p.add_run(text)
    set_run_font(run, bold=True)
    return p


def page_break(anchor: Paragraph) -> None:
    p = anchor.insert_paragraph_before()
    p.style = "Normal"
    p.add_run().add_break(WD_BREAK.PAGE)


def remove_body_between(document: Document, start_p, end_p) -> None:
    body = document._element.body
    children = list(body)
    start_index = children.index(start_p)
    end_index = children.index(end_p)
    for element in children[start_index + 1:end_index]:
        body.remove(element)


def find_paragraph(document: Document, text: str) -> Paragraph:
    normalized = text.casefold()
    for p in document.paragraphs:
        if p.text.strip().casefold() == normalized:
            return p
    raise RuntimeError(f"No se encontró el párrafo: {text}")


def section_break_before(document: Document, heading: Paragraph) -> Paragraph:
    body = document._element.body
    children = list(body)
    index = children.index(heading._p)
    for element in reversed(children[:index]):
        if element.tag == qn("w:p"):
            p = Paragraph(element, document._body)
            if p._p.pPr is not None and p._p.pPr.sectPr is not None:
                return p
    raise RuntimeError("No se encontró el salto de sección anterior.")


def create_current_process(path: Path) -> None:
    image = Image.new("RGB", (1800, 1100), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title = load_fonts(40, bold=True)
    draw.text((900, 48), "Proceso manual actual de compras y ventas", font=title, fill="#17365D", anchor="mm")
    top = ["Revisar estantes", "Pedir al proveedor", "Recibir productos", "Revisar comprobante", "Anotar compra y percepción", "Guardar documentos"]
    bottom = ["Atender venta", "Sumar importes", "Escribir ticket", "Cobrar o anotar fiado", "Ajustar caja de memoria", "Revisar cuadernos"]
    colors = ["FFF2CC", "FCE4D6"]
    for row, labels in enumerate((top, bottom)):
        y1 = 155 + row * 420
        boxes = []
        for i, label in enumerate(labels):
            x1 = 35 + i * 292
            box = (x1, y1, x1 + 245, y1 + 190)
            rounded_box(draw, box, label, colors[row], font_size=23)
            boxes.append(box)
            if i:
                arrow(draw, (boxes[i - 1][2] + 6, y1 + 95), (box[0] - 6, y1 + 95), width=5)
    warning = load_fonts(25, bold=True)
    draw.text((900, 500), "Registros separados: comprobantes, cuadernos, tickets y memoria", font=warning, fill="#9C0006", anchor="mm")
    small = load_fonts(23)
    draw.text((900, 1015), "El inventario, la caja, las percepciones y los fiados no se actualizan de forma integrada", font=small, fill="#404040", anchor="mm")
    image.save(path)


def create_ishikawa(path: Path) -> None:
    image = Image.new("RGB", (1800, 1100), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title = load_fonts(38, bold=True)
    draw.text((900, 46), "Causas de la deficiente gestión y control de la tienda", font=title, fill="#17365D", anchor="mm")
    spine_y = 560
    draw.line((160, spine_y, 1550, spine_y), fill="#35546E", width=10)
    arrow(draw, (1550, spine_y), (1710, spine_y), color="35546E", width=10)
    rounded_box(draw, (1450, 420, 1770, 700), "Errores, demoras y falta de información confiable", "F4CCCC", font_size=25)
    branches = [
        ("Métodos", ["Registros manuales", "Procesos no estandarizados"], 330, True),
        ("Personas", ["Sobrecarga de tareas", "Dependencia de la memoria"], 700, True),
        ("Información", ["Documentos dispersos", "Sin historial consolidado"], 1070, True),
        ("Tecnología", ["Sin sistema integrado", "Tickets manuscritos"], 420, False),
        ("Control", ["Sin alertas de percepciones", "Sin conciliación de caja"], 820, False),
        ("Entorno", ["Aglomeraciones", "Cámaras aisladas"], 1190, False),
    ]
    head = load_fonts(27, bold=True)
    body_font = load_fonts(22)
    for name, causes, x, upper in branches:
        if upper:
            draw.line((x, spine_y, x - 150, 220), fill="#5B6573", width=6)
            ty = 145
        else:
            draw.line((x, spine_y, x - 150, 900), fill="#5B6573", width=6)
            ty = 920
        draw.text((x - 150, ty), name, font=head, fill="#17365D", anchor="mm")
        draw.text((x - 150, ty + (42 if upper else 48)), "\n".join(causes), font=body_font, fill="#222222", anchor="ma", align="center")
    image.save(path)


def create_priorities(path: Path) -> None:
    image = Image.new("RGB", (1800, 1050), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title = load_fonts(40, bold=True)
    draw.text((900, 45), "Priorización de causas raíz", font=title, fill="#17365D", anchor="mm")
    labels = ["Sin sistema integral", "Compras y percepciones manuales", "Registros no estandarizados", "Stock no actualizado", "Tickets y cálculos manuales", "Fiados en cuaderno", "Sin historial de clientes", "Cámaras no integradas"]
    values = [15, 15, 14, 13, 12, 11, 9, 8]
    label_font = load_fonts(24)
    value_font = load_fonts(24, bold=True)
    for i, (label, value) in enumerate(zip(labels, values)):
        y = 125 + i * 108
        draw.text((35, y + 30), label, font=label_font, fill="#222222")
        x1 = 620
        x2 = x1 + value * 70
        color = "#C00000" if value >= 14 else "#ED7D31" if value >= 11 else "#5B9BD5"
        draw.rounded_rectangle((x1, y, x2, y + 65), radius=12, fill=color)
        draw.text((x2 + 18, y + 32), str(value), font=value_font, fill="#111111", anchor="lm")
    axis = load_fonts(20)
    draw.text((1665, 1010), "Puntaje máximo: 15", font=axis, fill="#555555", anchor="rm")
    image.save(path)


def crop_er_panels(left_path: Path, right_path: Path) -> None:
    image = Image.open(ER_SOURCE).convert("RGB")
    width, height = image.size
    overlap = int(width * 0.07)
    middle = width // 2
    image.crop((0, 0, middle + overlap, height)).save(left_path, quality=95)
    image.crop((middle - overlap, 0, width, height)).save(right_path, quality=95)


def insert_er_figure(anchor: Paragraph, left_path: Path, right_path: Path) -> None:
    insert_caption(anchor, "Figura 8", "Diagrama entidad–relación de Innova Store")
    p = anchor.insert_paragraph_before()
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("(a) Compras, proveedores, productos e inventario")
    set_run_font(r, bold=True, italic=True)
    pic = anchor.insert_paragraph_before()
    pic.style = "Normal"
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_together = True
    pic.add_run().add_picture(str(left_path), width=Inches(5.55))
    page_break(anchor)
    p = anchor.insert_paragraph_before()
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("(b) Ventas, caja, clientes, usuarios y seguridad")
    set_run_font(r, bold=True, italic=True)
    pic = anchor.insert_paragraph_before()
    pic.style = "Normal"
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_together = True
    pic.add_run().add_picture(str(right_path), width=Inches(5.55))
    insert_note(anchor, "Elaboración propia a partir del modelo diseñado en MySQL Workbench.")


def add_table_block(document: Document, anchor: Paragraph, number: int, title: str, headers, rows, widths, font_size=8.5) -> None:
    insert_caption(anchor, f"Tabla {number}", title)
    insert_table(document, anchor, headers, rows, widths, font_size=font_size)
    insert_note(anchor)


def add_chapter_three(document: Document, anchor: Paragraph, assets: dict[str, Path]) -> None:
    chapter_subtitle(anchor, "ANÁLISIS DE LA SITUACIÓN ACTUAL")
    insert_body(anchor, "El análisis de la situación actual se realizó en Minimarket Tienda Perez, negocio de propiedad de Perez Hancco Eudocio, ubicado en Ciudad Municipal, Zona 5, Manzana G, Lote 10, distrito de Cerro Colorado. La tienda comercializa productos de consumo cotidiano y desarrolla sus compras, ventas, control de existencias, caja, fiados y seguimiento de clientes mediante documentos físicos, cuadernos y cálculos manuales.")
    insert_body(anchor, "La revisión se concentra en la falta de integración de la información. El problema más crítico se presenta en el registro de compras y percepciones emitidas por proveedores, porque su control manual dificulta conocer el costo real, el acumulado mensual de adquisiciones y los importes que deben conservarse para una posterior revisión tributaria. Esta situación se relaciona con errores de inventario, tickets manuscritos, saldos de fiados desactualizados y reportes poco confiables.")

    section_heading(anchor, "3.1 Diagrama del proceso, mapa del flujo de valor y/o diagrama de operación actual")
    insert_body(anchor, "El proceso actual comienza con la observación de los estantes para decidir qué productos deben reponerse. Los pedidos se realizan a proveedores como empresas distribuidoras de bebidas y otros mayoristas. Al recibir la mercadería, los propietarios revisan las cantidades y conservan la factura, boleta, comprobante de percepción u otro documento entregado por el proveedor; sin embargo, la compra no queda vinculada automáticamente con el inventario, el egreso de caja ni el acumulado del periodo.")
    insert_body(anchor, "En la venta, los productos se seleccionan y suman manualmente. El ticket interno se escribe con papel y lapicero, y cuando la operación es al crédito el cliente, el monto y los abonos se anotan en un cuaderno. La información de clientes frecuentes depende principalmente de la memoria de los propietarios, mientras que las cámaras de seguridad se consultan por separado. La Figura 2 resume este flujo y evidencia los puntos donde se pierde continuidad de la información.")
    insert_figure(anchor, assets["current"], "Figura 2", "Proceso manual actual de compras, percepciones, ventas y control", width=6.15)
    add_table_block(document, anchor, 1, "Secuencia y riesgos del proceso manual actual", ["N.º", "Actividad", "Registro actual", "Riesgo o dificultad"], CURRENT_PROCESS_ROWS, [0.38, 1.48, 1.75, 2.55], 7.7)
    insert_body(anchor, "El proceso contiene actividades necesarias, pero también búsquedas, verificaciones repetidas y conciliaciones manuales que no agregan valor al cliente. La falta de una relación única entre compra, comprobante, percepción, pago y movimiento de inventario impide seguir la operación de principio a fin y obliga a reconstruirla a partir de varios documentos.")

    section_heading(anchor, "3.2 Efectos del problema en el área de trabajo o en los resultados de la empresa")
    insert_body(anchor, "Los efectos del problema abarcan la operación diaria y la administración del negocio. En compras, la falta de un registro estructurado de percepciones dificulta comprobar montos y periodos; en inventario, impide conocer las existencias reales; y en ventas, aumenta el riesgo de equivocarse en el total o de no conservar el detalle de lo vendido.")
    insert_body(anchor, "También se producen diferencias entre ventas, egresos y efectivo disponible, mientras que los fiados dependen de anotaciones que pueden quedar incompletas. Estas deficiencias reducen la capacidad de los propietarios para evaluar gastos, planificar reposiciones, reconocer a clientes frecuentes y tomar decisiones con información verificable.")
    add_table_block(document, anchor, 2, "Efectos de la gestión manual en el negocio", ["Área", "Situación observada", "Efecto principal", "Impacto"], EFFECT_ROWS, [1.18, 1.72, 2.72, 0.65], 8.0)
    insert_body(anchor, "La acumulación de estos efectos incrementa el tiempo dedicado a revisar cuadernos y documentos y limita el control preventivo. El negocio puede continuar atendiendo, pero no dispone de una visión consolidada y oportuna de sus compras, percepciones, stock, caja y cuentas por cobrar.")

    section_heading(anchor, "3.3 Análisis de las causas raíz que generan el problema")
    insert_body(anchor, "Para identificar las causas se emplearon el diagrama de Ishikawa y el método de los 5 Porqués. Las causas se agruparon en métodos, personas, información, tecnología, control y entorno. La Figura 3 muestra que los errores visibles no se originan únicamente en la digitación o en la edad de los usuarios, sino en la ausencia de procedimientos integrados, validaciones y una fuente única de información.")
    insert_figure(anchor, assets["ishikawa"], "Figura 3", "Diagrama de Ishikawa de la gestión deficiente de la tienda", width=6.15)
    add_table_block(document, anchor, 3, "Análisis de la causa principal mediante los 5 Porqués", ["Nivel", "Pregunta", "Respuesta"], WHY_ROWS, [0.55, 2.55, 3.15], 8.2)
    insert_body(anchor, "El análisis establece como causa raíz la ausencia de un sistema adaptado al negocio que relacione los procesos y aplique reglas consistentes. Los registros manuales son una manifestación de esa causa: al no existir una estructura común, los mismos datos se repiten, se omiten o no pueden conciliarse posteriormente.")

    section_heading(anchor, "3.4 Priorización de causas raíz")
    insert_body(anchor, "Las causas identificadas se priorizaron mediante factores cualitativos de frecuencia, impacto y urgencia. Cada criterio se valoró en una escala de 1 a 5 y el puntaje total se obtuvo mediante la suma de las tres valoraciones. Este procedimiento permite ordenar la intervención aun cuando el negocio no dispone de mediciones históricas completas.")
    add_table_block(document, anchor, 4, "Criterios de evaluación de las causas raíz", ["Criterio", "Descripción", "Escala"], CRITERIA_ROWS, [1.15, 3.45, 1.72], 8.5)
    add_table_block(document, anchor, 5, "Priorización cualitativa de las causas raíz", ["Causa", "F", "I", "U", "Total", "Prioridad"], PRIORITY_ROWS, [2.9, 0.42, 0.42, 0.42, 0.65, 1.0], 8.0)
    insert_figure(anchor, assets["priorities"], "Figura 4", "Priorización de causas mediante factores cualitativos", width=5.55)
    insert_subheading(anchor, "Conclusión del diagnóstico")
    insert_body(anchor, "Las causas de mayor prioridad son la ausencia de un sistema integral y el control manual de compras y percepciones, ambas con 15 puntos. Les siguen los registros no estandarizados, el inventario desactualizado, los cálculos manuales y los fiados en cuadernos. Por ello, la mejora debe centralizar primero compras, percepciones, inventario, ventas, caja y cuentas por cobrar; luego podrá ampliar clientes frecuentes y cámaras sin comprometer los procesos críticos.")


def add_chapter_four(document: Document, anchor: Paragraph, assets: dict[str, Path]) -> None:
    chapter_subtitle(anchor, "PROPUESTA TÉCNICA DE LA MEJORA")
    insert_body(anchor, "La propuesta consiste en implementar Innova Store como un sistema web para centralizar las operaciones de Minimarket Tienda Perez. La solución registrará compras, documentos emitidos por proveedores, percepciones, inventario, ventas, tickets internos, caja, fiados, clientes frecuentes y reportes. También permitirá visualizar cámaras autorizadas e incorporar alertas de apoyo mediante inteligencia artificial.")
    insert_body(anchor, "La prioridad de implementación será el control de compras y percepciones, seguido del inventario, las ventas, la caja y las cuentas por cobrar. El sistema no emitirá boletas ni facturas ni sustituirá la revisión contable o tributaria; organizará la información y generará alertas para facilitar su verificación.")

    section_heading(anchor, "4.1 Plan de acción de la mejora propuesta")
    insert_body(anchor, "El plan de acción organiza las actividades desde la validación de requisitos hasta el despliegue. Cada resultado será revisado con los propietarios mediante casos reales controlados, priorizando que la interfaz sea sencilla y que una operación pueda seguirse desde su registro hasta el reporte correspondiente.")
    add_table_block(document, anchor, 6, "Plan de acción para la implementación de Innova Store", ["N.º", "Acción", "Actividad principal", "Responsable", "Resultado"], PLAN_ROWS, [0.35, 1.22, 2.35, 1.27, 1.25], 7.4)
    insert_body(anchor, "La aceptación comprobará que una compra actualiza el inventario y registra sus documentos y percepciones; que una venta actualiza stock, caja o deuda; y que los reportes reproducen la información registrada. Las cámaras y el análisis automático se activarán después de estabilizar los módulos comerciales.")

    section_heading(anchor, "4.2 Consideraciones técnicas, operativas y ambientales para la implementación de la mejora")
    insert_subheading(anchor, "Consideraciones técnicas")
    insert_body(anchor, "La arquitectura será modular: Vue.js administrará la interfaz; NestJS aplicará las reglas del negocio; Prisma ORM gestionará el acceso a MySQL; y Socket.IO distribuirá eventos autorizados en tiempo real. Las confirmaciones de compras, ventas, pagos y abonos se ejecutarán mediante validaciones y transacciones para evitar actualizaciones parciales.")
    insert_body(anchor, "El módulo de percepciones conservará el vínculo con el proveedor, la compra y el comprobante recibido. Los umbrales y porcentajes se mantendrán como parámetros configurables, debido a que las condiciones tributarias pueden cambiar y requieren revisión con información oficial o asesoría especializada.")
    insert_subheading(anchor, "Consideraciones operativas")
    insert_body(anchor, "La carga inicial partirá de productos, proveedores, existencias, clientes, deudas y documentos que puedan verificarse. La interfaz empleará controles visibles, textos claros, confirmaciones y búsquedas sencillas para facilitar el uso por parte de los propietarios. El ticket generado será una constancia interna y mostrará expresamente que no constituye boleta ni factura.")
    insert_subheading(anchor, "Consideraciones ambientales y de seguridad")
    insert_body(anchor, "La digitalización reducirá cuadernos y copias internas; los tickets se imprimirán únicamente cuando sean necesarios. El servicio de video utilizará resoluciones y frecuencias acordes con el objetivo de seguridad, y toda alerta será tratada como una señal que debe ser revisada por una persona, no como confirmación automática de un hurto.")
    add_table_block(document, anchor, 7, "Matriz de consideraciones para la implementación", ["Aspecto", "Consideración", "Medida propuesta"], CONSIDERATION_ROWS, [1.25, 2.25, 2.95], 8.0)

    section_heading(anchor, "4.3 Recursos técnicos para implementar la mejora propuesta")
    insert_body(anchor, "Los recursos se definieron según la arquitectura y los módulos previstos. El servidor y el procesamiento de video se dimensionarán después de probar la cantidad y resolución de las cámaras. Se aprovecharán equipos compatibles ya disponibles para reducir el costo inicial.")
    add_table_block(document, anchor, 8, "Recursos de software", ["Componente", "Tecnologías", "Finalidad"], SOFTWARE_ROWS, [1.25, 2.65, 2.55], 8.0)
    add_table_block(document, anchor, 9, "Recursos de hardware e infraestructura", ["Recurso", "Especificación referencial", "Uso"], HARDWARE_ROWS, [1.4, 2.8, 2.25], 8.0)
    add_table_block(document, anchor, 10, "Recursos humanos y responsabilidades", ["Participante", "Responsabilidad"], HUMAN_ROWS, [2.0, 4.45], 8.5)
    insert_body(anchor, "También se requiere información validada: catálogo de productos y unidades, precios, proveedores, medios de pago, stock inicial, deudas, compras pendientes y comprobantes de percepción disponibles. Las copias de seguridad y los criterios de conservación de imágenes se acordarán antes del despliegue.")

    section_heading(anchor, "4.4 Diagrama del proceso, mapa del flujo de valor y diagrama de operación de la situación mejorada")
    insert_body(anchor, "La situación mejorada reemplaza los registros aislados por un flujo integrado. La compra relaciona proveedor, documento recibido, percepción, pago, productos y entrada de inventario. La venta relaciona detalle, ticket interno, medio de pago, salida de stock, caja, cliente y, cuando corresponda, deuda y abonos.")
    insert_figure(anchor, assets["improved"], "Figura 5", "Proceso mejorado de compras, ventas y control", width=6.15)
    insert_body(anchor, "El sistema actualizará el acumulado mensual de adquisiciones al confirmar cada compra y emitirá avisos según parámetros vigentes. La alerta informará que el valor requiere revisión; no decidirá por sí sola la categoría tributaria ni realizará declaraciones ante la SUNAT.")
    add_table_block(document, anchor, 11, "Secuencia de operación de la situación mejorada", ["N.º", "Actividad", "Control del sistema", "Resultado"], IMPROVED_PROCESS_ROWS, [0.4, 1.75, 2.55, 1.75], 7.8)
    insert_figure(anchor, assets["architecture"], "Figura 6", "Arquitectura funcional de Innova Store", width=6.15)
    insert_body(anchor, "El frontend concentra la interacción y el backend aplica las reglas del negocio. MySQL conserva la información transaccional; PDF y Excel permiten distribuir reportes; Socket.IO mantiene las pantallas sincronizadas; y el servicio de inteligencia artificial permanece desacoplado para que una falla de video no interrumpa compras, ventas o caja.")
    insert_figure(anchor, assets["ai"], "Figura 7", "Flujo de generación y revisión de alertas de seguridad", width=5.8)
    insert_body(anchor, "Las cámaras proporcionarán video mediante RTSP. OpenCV obtendrá los fotogramas; YOLOv8 detectará personas u objetos; ByteTrack mantendrá trayectorias; y el modelo temporal estimará patrones de interés. Las alertas llegarán mediante Socket.IO y serán revisadas por un usuario autorizado.")
    insert_er_figure(anchor, assets["er_left"], assets["er_right"])
    insert_body(anchor, "El modelo entidad–relación organiza la información en módulos relacionados. Compras, comprobantes, percepciones y egresos se vinculan con proveedores, productos, sucursales e inventario; ingresos, detalles de pago, caja y deudas se relacionan con clientes; y los usuarios se controlan mediante roles, permisos y sucursales. Esta estructura permite trazabilidad sin modificar el alcance tributario del ticket interno.")

    section_heading(anchor, "4.5 Control de compras y percepciones")
    insert_body(anchor, "En el Nuevo Régimen Único Simplificado existen dos categorías mensuales: la primera comprende ingresos o adquisiciones de hasta S/ 5 000 y la segunda, montos mayores a S/ 5 000 y hasta S/ 8 000. Además, se establece como límite general S/ 8 000 mensuales o S/ 96 000 anuales para ingresos o adquisiciones (SUNAT, s. f.-a). Por ello, el sistema mostrará el acumulado del periodo y alertará al aproximarse o superar los umbrales configurados.")
    insert_body(anchor, "Las percepciones del IGV pueden aplicarse contra la cuota mensual del Nuevo RUS; el saldo no utilizado puede arrastrarse a periodos posteriores o solicitarse en devolución cuando corresponda (SUNAT, s. f.-c). Innova Store almacenará los datos necesarios para preparar la revisión, pero no calculará obligaciones definitivas, presentará declaraciones ni reemplazará la validación del propietario o de un profesional contable.")
    add_table_block(document, anchor, 12, "Controles para el registro y seguimiento de percepciones", ["Dato o control", "Validación del sistema", "Registro generado", "Finalidad"], PERCEPTION_ROWS, [1.25, 2.2, 1.45, 1.55], 7.3)
    insert_body(anchor, "El reporte mensual permitirá contrastar las compras con sus documentos, percepciones y egresos. Las correcciones conservarán usuario, fecha, motivo y valores anteriores para evitar que un cambio elimine la trazabilidad del periodo.")

    section_heading(anchor, "4.6 Aspectos limitantes para la implementación de la mejora")
    insert_body(anchor, "La implementación presenta limitaciones técnicas, operativas y organizacionales. Los documentos históricos pueden estar incompletos, las reglas tributarias pueden cambiar y la conexión de las cámaras depende de compatibilidad y capacidad de procesamiento. Estas condiciones requieren parámetros editables, respaldos, pruebas graduales y revisión periódica.")
    add_table_block(document, anchor, 13, "Matriz de limitaciones, riesgos y medidas de respuesta", ["Riesgo", "Prob.", "Impacto", "Prevención", "Respuesta"], RISK_ROWS, [1.45, 0.55, 0.6, 2.0, 1.85], 7.2)
    insert_body(anchor, "La estrategia principal será poner en operación primero compras, comprobantes, percepciones, inventario, ventas, caja y fiados. Después se habilitarán clientes frecuentes, eventos en tiempo real y videovigilancia inteligente. La aceptación se sustentará en operaciones trazables, saldos coherentes, alertas verificables, permisos aplicados y respaldos recuperables.")


def add_reference(document: Document) -> None:
    target = None
    for p in document.paragraphs:
        if p.text.startswith("Ultralytics."):
            target = p
            break
    if target is None:
        raise RuntimeError("No se encontró el punto de inserción de referencias.")
    text = ("Superintendencia Nacional de Aduanas y de Administración Tributaria. (s. f.-c). "
            "Instructivo para hacer el pago del Nuevo RUS con percepciones del IGV y otros medios de pago. "
            "Recuperado el 28 de septiembre de 2026, de "
            "https://orientacion.sunat.gob.pe/sites/default/files/inline-files/nrus_percepciones_ultimo_compressed_0.pdf")
    p = target.insert_paragraph_before()
    p.style = target.style
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not ER_SOURCE.exists():
        raise FileNotFoundError(ER_SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    assets = {
        "current": ASSET_DIR / "chapter3-current-process.png",
        "ishikawa": ASSET_DIR / "chapter3-ishikawa.png",
        "priorities": ASSET_DIR / "chapter3-priorities.png",
        "improved": ASSET_DIR / "chapter4-process.png",
        "architecture": ASSET_DIR / "chapter4-architecture.png",
        "ai": ASSET_DIR / "chapter4-ai-flow.png",
        "er_left": ASSET_DIR / "er-left.png",
        "er_right": ASSET_DIR / "er-right.png",
    }
    create_current_process(assets["current"])
    create_ishikawa(assets["ishikawa"])
    create_priorities(assets["priorities"])
    create_process_diagram(assets["improved"])
    create_architecture_diagram(assets["architecture"])
    create_ai_diagram(assets["ai"])
    crop_er_panels(assets["er_left"], assets["er_right"])

    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    chapter_three = find_paragraph(document, "CAPÍTULO III")
    chapter_four = find_paragraph(document, "CAPÍTULO IV")
    try:
        chapter_five = find_paragraph(document, "CAPITULO V")
    except RuntimeError:
        chapter_five = find_paragraph(document, "CAPÍTULO V")
    break_before_four = section_break_before(document, chapter_four)
    break_before_five = section_break_before(document, chapter_five)

    remove_body_between(document, chapter_four._p, break_before_five._p)
    add_chapter_four(document, break_before_five, assets)
    remove_body_between(document, chapter_three._p, break_before_four._p)
    add_chapter_three(document, break_before_four, assets)
    add_reference(document)

    document.core_properties.title = "Sistema web integral para la gestión de Minimarket Tienda Perez"
    document.core_properties.subject = "Capítulos III y IV actualizados: diagnóstico, propuesta, percepciones y modelo entidad–relación"
    document.save(OUTPUT)
    print(f"Documento actualizado: {OUTPUT}")
    print(f"Párrafos: {len(document.paragraphs)} | Tablas: {len(document.tables)} | Imágenes: {len(document.inline_shapes)}")


if __name__ == "__main__":
    main()
